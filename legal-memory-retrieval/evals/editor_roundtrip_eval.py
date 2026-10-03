"""Document editor round-trip fidelity and latency, through the running API (plan 16, E5).

For contract-style Word files of increasing length (the plan-15 long-document generator):
upload as a version → read the edit model → lock → save a batch of paragraph edits as a
tracked version → download the stored file, and check:

  read_ok        the edit model's paragraphs are exactly the file's body paragraphs
  accept_ok      the file with all changes accepted equals the expected edited text
  reject_ok      the file with every change rejected reads as the document first uploaded
                 (saves keep earlier pending changes, including the editor's own)
  untouched_ok   every paragraph not edited keeps its style and run formatting (100%)
  tables_ok      tables (view-only in the editor) survive unchanged
  next_ok        the next edit model starts from the accepted text (a second round runs on it)
  format_accept_ok  formatting edits (bold a phrase, restyle a paragraph) read back exactly
                 from the saved file with all changes accepted
  format_reject_ok  rejecting the formatting revisions restores those paragraphs' original
                 style and run formatting
  tracked_render_ok the saved tracked file converts in LibreOffice (Gotenberg), an
                 independent check that the Word XML is valid
  vectors_ok     every chunk of the new version gets its vector (unchanged chunks reuse the
                 previous one at save time; the rest are embedded in the background)

and times save, edit-model load and the exact (PDF) rendition.

Gates: all checks 100%; save p95 < 3 s on the 400-page document.

    python evals/editor_roundtrip_eval.py --base-url http://localhost:8011 [--pages 10,100,400]

Run the API without --reload. The throwaway documents are deleted at the end.
"""
from __future__ import annotations

import argparse
import io
import json
import random
import statistics
import sys
import time
from pathlib import Path

import httpx
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.db.connection import connect  # noqa: E402
from app.documents.docx_review import accept_everything, read_revisions, resolve, text_view  # noqa: E402
from app.drafting.docx_tracked import formatting_signature  # noqa: E402
from evals.long_doc.generate import build_document, to_docx  # noqa: E402

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
OUT = ROOT / "evals" / "last_editor_roundtrip.json"
SAVE_P95_GATE_S = 3.0


def _pick_editor() -> tuple[str, str]:
    """An open matter and a non-admin member of its team."""
    with connect() as conn:
        admins = {r["member_id"] for r in conn.execute(
            "SELECT DISTINCT member_id FROM member_roles WHERE role_key IN ('firm_admin', 'risk_compliance')")}
        for m in conn.execute(
            """SELECT m.matter_id FROM matters m JOIN matter_access a USING (matter_id)
               WHERE a.mode = 'open' ORDER BY m.matter_id""").fetchall():
            for r in conn.execute("SELECT member_id FROM matter_members WHERE matter_id = %s ORDER BY member_id",
                                  (m["matter_id"],)):
                if r["member_id"] not in admins:
                    return m["matter_id"], r["member_id"]
    raise SystemExit("no open matter with a non-admin team member")


def _cleanup(doc_ids: list[str]) -> None:
    with connect() as conn:
        for doc_id in doc_ids:
            for table in ("document_events", "document_drafts", "document_locks", "chunks", "document_blocks", "version_diffs"):
                conn.execute(f"DELETE FROM {table} WHERE document_id = %s", (doc_id,))
            conn.execute("UPDATE documents SET current_version_id = NULL WHERE document_id = %s", (doc_id,))
            conn.execute("DELETE FROM document_versions WHERE document_id = %s", (doc_id,))
            conn.execute("DELETE FROM documents WHERE document_id = %s", (doc_id,))
        conn.commit()


def _body_paragraphs(data: bytes) -> list:
    """Body paragraph elements outside tables, in order (what the editor's pids index)."""
    body = Document(io.BytesIO(data)).element.body
    return [p for p in body.iter(qn("w:p")) if not any(a.tag == qn("w:tbl") for a in p.iterancestors())]


def _tables(data: bytes) -> list[list[list[str]]]:
    return [[[c.text for c in row.cells] for row in t.rows] for t in Document(io.BytesIO(data)).tables]


def _make_ops(texts: list[str], rng: random.Random, n_replace: int, n_delete: int, n_insert: int) -> list[dict]:
    """Realistic edits on non-empty paragraphs: in-sentence word changes, appended sentences,
    deletions and new clauses."""
    candidates = [i for i, t in enumerate(texts) if len(t) > 40]
    rng.shuffle(candidates)
    ops: list[dict] = []
    for pid in candidates[:n_replace]:
        words = texts[pid].split(" ")
        if rng.random() < 0.5 and len(words) > 6:
            k = rng.randrange(2, len(words) - 1)
            words[k] = "reasonable" if words[k] != "reasonable" else "prompt"
            new = " ".join(words)
        else:
            new = texts[pid] + " This obligation survives termination."
        ops.append({"op": "replace", "pid": pid, "text": new})
    for pid in candidates[n_replace:n_replace + n_delete]:
        ops.append({"op": "delete", "pid": pid})
    for pid in candidates[n_replace + n_delete:n_replace + n_delete + n_insert]:
        ops.append({"op": "insert_after", "pid": pid, "text": f"{pid}.A The Parties shall co-operate in good faith."})
    return ops


def _accept(data: bytes) -> bytes:
    return accept_everything(data)


def _formatting_rejected(data: bytes) -> bytes:
    """Formatting revisions rejected, every other change accepted."""
    fmt = [r.key for r in read_revisions(data) if r.type in ("format", "paragraph_format")]
    out, _ = resolve(data, fmt, accept=False)
    return accept_everything(out)


def _format_ops(paragraphs: list[dict], used: set[int], rng: random.Random, n: int) -> list[dict]:
    """Formatting-only edits on paragraphs no other op touches: bold a phrase inside the
    paragraph (keeping its other formatting), or restyle it as Heading 2."""
    pool = [p for p in paragraphs if p["pid"] not in used and len(p["text"]) > 40
            and "".join(r["text"] for r in p["runs"]) == p["text"]]
    rng.shuffle(pool)
    ops = []
    for i, p in enumerate(pool[:n]):
        if i % 2:
            ops.append({"op": "format", "pid": p["pid"], "style": "Heading 2"})
            continue
        chars = [(c, r["bold"], r["italic"], r["underline"]) for r in p["runs"] for c in r["text"]]
        a = len(chars) // 3
        b = min(len(chars), a + 12)
        chars = [(c, True if a <= k < b else bold, it, un) for k, (c, bold, it, un) in enumerate(chars)]
        runs: list[dict] = []
        for c, bold, it, un in chars:
            if runs and (runs[-1]["bold"], runs[-1]["italic"], runs[-1]["underline"]) == (bold, it, un):
                runs[-1]["text"] += c
            else:
                runs.append({"text": c, "bold": bold, "italic": it, "underline": un})
        ops.append({"op": "format", "pid": p["pid"], "runs": runs})
    return ops


def _char_flags(p_el) -> list[tuple[bool, bool, bool]]:
    from docx.text.paragraph import Paragraph

    return [(bool(r.bold), bool(r.italic), bool(r.underline)) for r in Paragraph(p_el, None).runs for _ in r.text]


def _format_checks(base: bytes, out: bytes, ops: list[dict], origin: list[int | None],
                   base_idx: dict[int, int]) -> tuple[bool, bool, int]:
    """(accepted formatting as asked, rejected formatting as before, formatting ops checked)."""
    fmt = [o for o in ops if o["op"] == "format"]
    if not fmt:
        return True, True, 0
    accepted = _body_paragraphs(_accept(out))
    position = {pid: i for i, pid in enumerate(origin) if pid is not None}
    accept_ok = len(accepted) == len(origin)
    for o in fmt:
        el = accepted[position[o["pid"]]] if accept_ok else None
        if el is None:
            break
        if o.get("style") is not None:
            accept_ok &= _style_name(out, el) == o["style"]
        if o.get("runs") is not None:
            want = [(bool(r["bold"]), bool(r["italic"]), bool(r["underline"])) for r in o["runs"] for _ in r["text"]]
            accept_ok &= _char_flags(el) == want
    old = _body_paragraphs(base)
    rejected = _body_paragraphs(_formatting_rejected(out))
    # With text changes accepted and formatting rejected, paragraphs line up with ``origin``.
    reject_ok = len(rejected) == len(origin) and all(
        formatting_signature(rejected[position[o["pid"]]]) == formatting_signature(old[base_idx[o["pid"]]]) for o in fmt)
    return bool(accept_ok), reject_ok, len(fmt)


def _style_name(data: bytes, p_el) -> str:
    from app.documents.docx_format import paragraph_styles

    ppr = p_el.find(qn("w:pPr"))
    sid = ppr.find(qn("w:pStyle")).get(qn("w:val")) if ppr is not None and ppr.find(qn("w:pStyle")) is not None else None
    names = {v: k for k, v in paragraph_styles(Document(io.BytesIO(data))).items()}
    return names.get(sid, "Normal") if sid else "Normal"


def _renders(data: bytes) -> bool:
    from app.documents.pdf_render import _gotenberg

    pdf = _gotenberg(data, ".docx")
    return bool(pdf) and pdf[:4] == b"%PDF"


def _expected(texts: list[str], ops: list[dict]) -> tuple[list[str], list[int | None]]:
    """Accepted text after the ops, and for each resulting paragraph the original pid (None = inserted)."""
    replaced = {o["pid"]: o["text"] for o in ops if o["op"] == "replace"}
    deleted = {o["pid"] for o in ops if o["op"] == "delete"}
    inserts: dict[int, list[str]] = {}
    for o in ops:
        if o["op"] == "insert_after":
            inserts.setdefault(o["pid"], []).append(o["text"])
    text_out, origin = [], []
    for pid, t in enumerate(texts):
        if pid not in deleted:
            text_out.append(replaced.get(pid, t))
            origin.append(pid)
        for new in inserts.get(pid, []):
            text_out.append(new)
            origin.append(None)
    return text_out, origin


def _vector_counts(version_id: str) -> tuple[int, int]:
    """(chunks with a vector, all chunks) of a version."""
    with connect() as conn:
        row = conn.execute("SELECT count(embedding) AS done, count(*) AS total FROM chunks WHERE version_id = %s",
                           (version_id,)).fetchone()
    return row["done"], row["total"]


def _wait_vectors(version_id: str, timeout: float = 180.0) -> float | None:
    """Seconds until every chunk of the version has its vector (embedded in the background), or None."""
    t0 = time.perf_counter()
    while time.perf_counter() - t0 < timeout:
        done, total = _vector_counts(version_id)
        if total and done == total:
            return round(time.perf_counter() - t0, 2)
        time.sleep(0.25)
    return None


def _p95(xs: list[float]) -> float:
    xs = sorted(xs)
    return xs[min(len(xs) - 1, int(round(0.95 * (len(xs) - 1))))] if xs else 0.0


def run(base_url: str, pages_list: list[int], rounds: int, seed: int) -> dict:
    matter_id, member = _pick_editor()
    headers = {"X-Member-Id": member}
    rng = random.Random(seed)
    created: list[str] = []
    results = []
    with httpx.Client(base_url=base_url, headers=headers, timeout=600) as http:
        try:
            for pages in pages_list:
                built = build_document(pages, seed)
                buf = io.BytesIO()
                to_docx(built, buf)
                original = buf.getvalue()

                r = http.post("/api/documents/ingest", json={
                    "title": f"Editor round-trip {pages}p", "matter_id": matter_id, "body": "placeholder",
                    "document_type": "Agreement"})
                r.raise_for_status()
                doc_id = r.json()["document_id"]
                created.append(doc_id)
                t0 = time.perf_counter()
                r = http.post(f"/api/editor/documents/{doc_id}/versions",
                              files={"file": (f"roundtrip-{pages}.docx", original, DOCX_MIME)}, data={"note": "eval"})
                r.raise_for_status()
                upload_s = time.perf_counter() - t0

                t0 = time.perf_counter()
                r = http.get(f"/api/documents/{doc_id}/render")
                render_s = time.perf_counter() - t0
                render_ok = r.status_code == 200 and r.content[:4] == b"%PDF"

                row = {"pages": pages, "paragraphs": 0, "upload_s": round(upload_s, 2),
                       "render_s": round(render_s, 2), "render_ok": render_ok, "rounds": []}
                current = original
                first_original = text_view(original, "original")
                for rnd in range(rounds):
                    t0 = time.perf_counter()
                    r = http.get(f"/api/editor/documents/{doc_id}")
                    r.raise_for_status()
                    model_s = time.perf_counter() - t0
                    model = r.json()
                    texts = [p["text"] for p in model["paragraphs"]]
                    row["paragraphs"] = row["paragraphs"] or len(texts)
                    base_accepted = _accept(current)
                    # Paragraphs deleted in an earlier round stay in the file (pending) and in the model
                    # (locked, empty) — they are gone from the Final view.
                    gone = {p["pid"] for p in model["paragraphs"] if p.get("locked_reason") == "deleted"}
                    live_pids = [p["pid"] for p in model["paragraphs"] if p["pid"] not in gone]
                    base_idx = {pid: i for i, pid in enumerate(live_pids)}
                    read_ok = [texts[pid] for pid in live_pids] == text_view(current, "final")

                    n = len(texts)
                    ops = _make_ops(texts, rng, n_replace=max(3, n // 40), n_delete=max(1, n // 200), n_insert=max(1, n // 200))
                    used = {o["pid"] for o in ops}
                    ops += _format_ops(model["paragraphs"], used, rng, max(2, n // 100))
                    http.post(f"/api/editor/documents/{doc_id}/lock").raise_for_status()
                    t0 = time.perf_counter()
                    r = http.post(f"/api/editor/documents/{doc_id}/save", json={
                        "base_version_id": model["base_version_id"], "ops": ops, "note": f"round {rnd + 1}", "mode": "tracked"})
                    save_s = time.perf_counter() - t0
                    http.delete(f"/api/editor/documents/{doc_id}/lock")
                    r.raise_for_status()
                    reused, total = _vector_counts(r.json()["version_id"])
                    vectors_s = _wait_vectors(r.json()["version_id"])

                    stored = http.get(f"/api/documents/{doc_id}/download")
                    stored.raise_for_status()
                    out = stored.content
                    want_all, origin_all = _expected(texts, ops)
                    kept = [(t, o) for t, o in zip(want_all, origin_all) if o is None or o not in gone]
                    want, origin = [t for t, _ in kept], [o for _, o in kept]
                    accepted = _accept(out)
                    accept_ok = text_view(out, "final") == want
                    reject_ok = text_view(out, "original") == first_original

                    touched = {o["pid"] for o in ops if o["op"] in ("replace", "delete", "format")}
                    format_accept_ok, format_reject_ok, format_checked = _format_checks(base_accepted, out, ops, origin, base_idx)
                    tracked_render_ok = _renders(out) if rnd == 0 else True
                    old_p = _body_paragraphs(base_accepted)
                    new_p = _body_paragraphs(accepted)
                    checked = kept = 0
                    if len(new_p) == len(origin):
                        for el, pid in zip(new_p, origin):
                            if pid is None or pid in touched:
                                continue
                            checked += 1
                            kept += formatting_signature(el) == formatting_signature(old_p[base_idx[pid]])
                    untouched_ok = checked > 0 and kept == checked
                    tables_ok = _tables(accepted) == _tables(base_accepted)

                    after = http.get(f"/api/editor/documents/{doc_id}").json()
                    next_ok = [p["text"] for p in after["paragraphs"] if p.get("locked_reason") != "deleted"] == want and not after["draft"]
                    row["rounds"].append({
                        "ops": len(ops), "stats": r.json()["stats"], "save_s": round(save_s, 3), "model_s": round(model_s, 3),
                        "read_ok": read_ok, "accept_ok": accept_ok, "reject_ok": reject_ok,
                        "untouched_checked": checked, "untouched_kept": kept, "untouched_ok": untouched_ok,
                        "tables_ok": tables_ok, "next_ok": next_ok,
                        "format_ops": format_checked, "format_accept_ok": format_accept_ok,
                        "format_reject_ok": format_reject_ok, "tracked_render_ok": tracked_render_ok,
                        "chunks": total, "vectors_reused_at_save": reused,
                        "vectors_complete_s": vectors_s, "vectors_ok": vectors_s is not None,
                    })
                    current = out
                saves = [x["save_s"] for x in row["rounds"]]
                row["save_p95_s"] = round(_p95(saves), 3)
                row["all_ok"] = render_ok and all(
                    x[k] for x in row["rounds"] for k in ("read_ok", "accept_ok", "reject_ok", "untouched_ok", "tables_ok", "next_ok",
                                                         "format_accept_ok", "format_reject_ok", "tracked_render_ok",
                                                         "vectors_ok"))
                results.append(row)
                print(f"{pages:>4}p  paras={row['paragraphs']:>5}  upload={upload_s:5.2f}s  render={render_s:5.2f}s  "
                      f"save={','.join(f'{s:.2f}' for s in saves)}s  model={','.join(str(x['model_s']) for x in row['rounds'])}s  "
                      f"reused={','.join(f"{x['vectors_reused_at_save']}/{x['chunks']}" for x in row['rounds'])}  "
                      f"vectors={','.join(str(x['vectors_complete_s']) for x in row['rounds'])}s  ok={row['all_ok']}", flush=True)
                if not row["all_ok"]:
                    print(json.dumps(row["rounds"], indent=1), flush=True)
        finally:
            _cleanup(created)

    largest = max(results, key=lambda x: x["pages"]) if results else None
    summary = {
        "all_checks_ok": all(x["all_ok"] for x in results),
        "largest_pages": largest["pages"] if largest else None,
        "largest_save_p95_s": largest["save_p95_s"] if largest else None,
        "gate_save_p95_s": SAVE_P95_GATE_S,
    }
    summary["pass"] = summary["all_checks_ok"] and (largest is None or largest["save_p95_s"] < SAVE_P95_GATE_S)
    return {"base_url": base_url, "member": member, "seed": seed, "summary": summary, "results": results}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--base-url", default="http://localhost:8011")
    ap.add_argument("--pages", default="10,100,400")
    ap.add_argument("--rounds", type=int, default=3)
    ap.add_argument("--seed", type=int, default=16)
    args = ap.parse_args()
    report = run(args.base_url, [int(x) for x in args.pages.split(",")], args.rounds, args.seed)
    OUT.write_text(json.dumps(report, indent=1))
    print(json.dumps(report["summary"], indent=1))
    sys.exit(0 if report["summary"]["pass"] else 1)


if __name__ == "__main__":
    main()
