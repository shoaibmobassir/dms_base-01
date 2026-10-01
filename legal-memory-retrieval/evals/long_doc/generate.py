"""Synthetic long contracts and editing tasks with exact gold results (plan 15, Part D1).

A document is a list of paragraphs, each with a style and formatted runs, rendered to .docx
with python-docx and to the paged text the Assistant reads ("[Page N]" markers roughly every
3,000 characters). Everything is generated from a seed, so a task's gold result — the full
paragraph list after the edit — is known exactly.

Traps are planted on purpose:
  - payment periods written several ways ("thirty (30) days", "30 days after receipt") next to
    other 30-day periods that are not payment terms (notice, cure, retention) and must not change;
  - references to Clause 9 next to references to Clause 19 and 29 (a naive find/replace breaks);
  - the defined term "Supplier" also inside words that must not change ("Suppliers' Forum" is a
    proper name kept verbatim by the rename task) — see ``rename_term``.
"""
from __future__ import annotations

import random
import re
from dataclasses import dataclass, field

CHARS_PER_PAGE = 3000

TOPICS = [
    "Services", "Service Levels", "Personnel", "Customer Obligations", "Change Control", "Fees and Payment",
    "Invoicing", "Term", "Termination", "Confidentiality", "Intellectual Property", "Data Protection",
    "Security", "Audit", "Warranties", "Indemnities", "Limitation of Liability", "Insurance",
    "Force Majeure", "Subcontracting", "Assignment", "Dispute Resolution", "Notices", "Governing Law",
]
FILLERS = [
    "The {S} shall perform its obligations under this clause with reasonable skill, care and diligence and in accordance with Good Industry Practice.",
    "Each party shall bear its own costs in connection with the matters described in this clause, save as otherwise expressly provided.",
    "The {C} may, acting reasonably, request such information from the {S} as it requires to verify compliance with this clause.",
    "Nothing in this clause shall limit any other right or remedy of the {C} under this Agreement or at law.",
    "The {S} shall maintain complete and accurate records relating to the {topic} and shall make them available to the {C} on request.",
    "Any approval given by the {C} under this clause shall not relieve the {S} of any of its obligations under this Agreement.",
    "The parties shall meet at least quarterly to review performance in relation to the {topic} and agree any improvements required.",
    "Where the {S} becomes aware of any matter which may affect the {topic}, it shall notify the {C} promptly and in any event within five Working Days.",
]
PAYMENT_FORMS = [
    "The {C} shall pay each valid invoice within thirty (30) days of receipt.",
    "Undisputed amounts shall be paid no later than 30 days after the date of the relevant invoice.",
    "Payment of the Service Credits reconciliation shall be made within thirty (30) days following the end of each Contract Year.",
    "The {C} shall settle the Milestone Payment within 30 days of the Acceptance Certificate.",
]
PAYMENT_GOLD = [
    "The {C} shall pay each valid invoice within forty-five (45) days of receipt.",
    "Undisputed amounts shall be paid no later than 45 days after the date of the relevant invoice.",
    "Payment of the Service Credits reconciliation shall be made within forty-five (45) days following the end of each Contract Year.",
    "The {C} shall settle the Milestone Payment within 45 days of the Acceptance Certificate.",
]
DISTRACTORS_30 = [
    "Either party may terminate this Agreement for convenience on not less than thirty (30) days' written notice.",
    "The {S} shall remedy any material breach within 30 days of written notice requiring it to do so.",
    "The {S} shall retain all Customer Data for 30 days after termination before secure deletion.",
    "The {S} shall provide the exit plan within thirty (30) days of the Effective Date.",
]


@dataclass
class Para:
    text: str
    style: str = "Normal"                 # Title | Heading 1 | Heading 2 | Normal | Table
    runs: list[tuple[str, bool, bool]] = field(default_factory=list)  # (text, bold, italic); "" → one plain run
    table: list[list[str]] | None = None  # for style == "Table": rows of cells; ``text`` is the flattened text

    def formatted_runs(self) -> list[tuple[str, bool, bool]]:
        return self.runs or [(self.text, False, False)]


@dataclass
class Doc:
    name: str
    paras: list[Para]
    planted: dict

    def texts(self) -> list[str]:
        return [p.text for p in self.paras]


def _clause(num: str, body: str, bold_terms: tuple[str, ...] = ()) -> Para:
    text = f"{num} {body}"
    runs: list[tuple[str, bool, bool]] = [(f"{num} ", True, False)]
    rest = body
    for term in bold_terms:
        i = rest.find(f"“{term}”")
        if i >= 0:
            runs += [(rest[:i], False, False), (f"“{term}”", True, False)]
            rest = rest[i + len(term) + 2:]
    if " Good Industry Practice" in rest and not bold_terms:
        i = rest.find("Good Industry Practice")
        runs += [(rest[:i], False, False), ("Good Industry Practice", False, True)]
        rest = rest[i + len("Good Industry Practice"):]
    runs.append((rest, False, False))
    return Para(text=text, runs=[r for r in runs if r[0]])


def build_document(pages: int, seed: int) -> Doc:
    rng = random.Random(seed)
    S, C = "Supplier", "Customer"
    paras: list[Para] = [Para("MASTER SERVICES AGREEMENT", "Title", [("MASTER SERVICES AGREEMENT", True, False)])]
    planted: dict = {"payment": [], "distractor30": [], "clause9_refs": [], "clause_x9_refs": []}
    paras.append(Para("1. Definitions", "Heading 1"))
    defs = [("Supplier", "means Northwind Services Limited and its permitted successors."),
            ("Customer", "means Harbour Holdings plc."),
            ("Working Day", "means a day other than a Saturday, Sunday or public holiday in London."),
            ("Good Industry Practice", "means the exercise of the skill and care reasonably expected of a leading supplier.")]
    for i, (term, meaning) in enumerate(defs, 1):
        paras.append(_clause(f"1.{i}", f"“{term}” {meaning}", (term,)))
    paras.append(Para("1.5 References to the Suppliers' Forum are to the industry body of that name.", runs=[
        ("1.5 ", True, False), ("References to the Suppliers' Forum are to the industry body of that name.", False, False)]))

    clause = 1
    target_chars = pages * CHARS_PER_PAGE
    used = sum(len(p.text) for p in paras)
    topic_cycle = 0
    while used < target_chars * 1.08:
        clause += 1
        topic = TOPICS[topic_cycle % len(TOPICS)] + ("" if topic_cycle < len(TOPICS) else f" (Module {topic_cycle // len(TOPICS) + 1})")
        topic_cycle += 1
        paras.append(Para(f"{clause}. {topic}", "Heading 1"))
        for sub in range(1, rng.randint(9, 14)):
            num = f"{clause}.{sub}"
            roll = rng.random()
            if topic.startswith("Fees and Payment") and sub in (2, 5) or (roll < 0.03 and len(planted["payment"]) < 12):
                form = rng.randrange(len(PAYMENT_FORMS))
                body = PAYMENT_FORMS[form].format(C=C, S=S)
                paras.append(_clause(num, body))
                planted["payment"].append((len(paras) - 1, form))
            elif roll < 0.07:
                body = rng.choice(DISTRACTORS_30).format(C=C, S=S)
                paras.append(_clause(num, body))
                planted["distractor30"].append(len(paras) - 1)
            elif roll < 0.11:
                ref = rng.choice(["9", "9.2", "9.4"])
                body = f"Subject to Clause {ref}, the {S} shall provide the deliverables described in this clause to the {C}."
                paras.append(_clause(num, body))
                planted["clause9_refs"].append(len(paras) - 1)
            elif roll < 0.14:
                ref = rng.choice(["19", "29.1", "19.3"])
                body = f"Clause {ref} applies to any dispute arising under this clause."
                paras.append(_clause(num, body))
                planted["clause_x9_refs"].append(len(paras) - 1)
            else:
                n_sent = rng.randint(2, 4)
                body = " ".join(rng.choice(FILLERS).format(S=S, C=C, n=num, topic=topic.lower()) for _ in range(n_sent))
                paras.append(_clause(num, body))
        used = sum(len(p.text) for p in paras)
    planted["last_clause"] = clause

    paras.append(Para("Schedule 1 — Service Description", "Heading 1"))
    for i in range(1, 5):
        paras.append(Para(f"S1.{i} The {S} shall deliver service line {i} in accordance with the Service Levels."))
    paras.append(Para("Schedule 2 — Charges", "Heading 1"))
    rows = [["Service line", "Monthly charge (GBP)", "Payment terms"]] + [[f"Line {i}", f"{12000 + 1500 * i:,}", "30 days"] for i in range(1, 5)]
    paras.append(Para(" | ".join(" ; ".join(r) for r in rows), "Table", table=rows))
    planted["schedule3_start"] = len(paras)
    paras.append(Para("Schedule 3 — Service Credits", "Heading 1"))
    for i in range(1, 6):
        paras.append(Para(f"S3.{i} Where Service Level {i} is missed in a month, a Service Credit of {i}% of the monthly charge applies."))
    planted["schedule3_end"] = len(paras)
    paras.append(Para("Schedule 4 — Exit Management", "Heading 1"))
    for i in range(1, 4):
        paras.append(Para(f"S4.{i} On expiry the {S} shall co-operate with any replacement supplier."))
    return Doc(name=f"msa_{pages}p_s{seed}", paras=paras, planted=planted)


# ---------------------------------------------------------------------------
# Paged text (what the Assistant reads) and .docx
# ---------------------------------------------------------------------------

def paged_text(texts: list[str]) -> str:
    out, page, size = ["[Page 1]"], 1, 0
    for t in texts:
        if size > CHARS_PER_PAGE:
            page += 1
            out.append(f"[Page {page}]")
            size = 0
        out.append(t)
        size += len(t)
    return "\n\n".join(out)


def to_docx(doc: Doc, path) -> None:
    from docx import Document

    d = Document()
    for p in doc.paras:
        if p.style == "Table" and p.table:
            t = d.add_table(rows=len(p.table), cols=len(p.table[0]))
            for r, row in enumerate(p.table):
                for c, cell in enumerate(row):
                    t.cell(r, c).text = cell
            continue
        para = d.add_paragraph(style=p.style if p.style != "Normal" else None)
        for text, bold, italic in p.formatted_runs():
            run = para.add_run(text)
            run.bold, run.italic = bold or None, italic or None
    d.save(path if hasattr(path, "write") else str(path))


# ---------------------------------------------------------------------------
# Tasks: instruction + exact gold paragraph list
# ---------------------------------------------------------------------------

_SUPPLIER = re.compile(r"\bSupplier\b(?!s' Forum)")


def rename_term(texts: list[str]) -> list[str]:
    return [_SUPPLIER.sub("Vendor", t) for t in texts]


def tasks_for(doc: Doc, rng: random.Random) -> list[dict]:
    texts = doc.texts()
    clause_paras = [i for i, t in enumerate(texts) if re.match(r"^\d+\.\d+ ", t) and not t.startswith("1.")]
    out: list[dict] = []

    out.append({"kind": "rename_term", "instruction": (
        "Rename the defined term “Supplier” to “Vendor” everywhere in this agreement, including the definition. "
        "Do not change the name of the industry body “Suppliers' Forum”."), "gold": rename_term(texts)})

    i = rng.choice([c for c in clause_paras if c not in doc.planted["payment"] + doc.planted["distractor30"]])
    num = texts[i].split(" ", 1)[0]
    gold = list(texts)
    gold[i] = texts[i].replace("reasonable", "all reasonable", 1) if "reasonable" in texts[i] else texts[i] + " This obligation is of the essence."
    change = "replace the word “reasonable” (its first occurrence) with “all reasonable”" if "reasonable" in texts[i] else "add at the end the sentence “This obligation is of the essence.”"
    out.append({"kind": "amend_clause", "instruction": f"In clause {num}, {change}. Change nothing else.", "gold": gold})

    j = rng.choice(clause_paras)
    num = texts[j].split(" ", 1)[0]
    major, minor = num.split(".")
    new_num = f"{major}.{minor}A"
    new_text = f"{new_num} The Vendor's obligations under this clause survive any novation of this Agreement."
    gold = list(texts)
    gold.insert(j + 1, new_text)
    out.append({"kind": "insert_after", "instruction": (
        f"Insert a new clause immediately after clause {num}, worded exactly: “{new_text}”. Change nothing else."), "gold": gold})

    s, e = doc.planted["schedule3_start"], doc.planted["schedule3_end"]
    out.append({"kind": "delete_schedule", "instruction": "Delete Schedule 3 (Service Credits) entirely, heading included. Change nothing else.",
                "gold": texts[:s] + texts[e:]})

    gold = list(texts)
    for idx, form in doc.planted["payment"]:
        gold[idx] = gold[idx].split(" ", 1)[0] + " " + PAYMENT_GOLD[form].format(C="Customer", S="Supplier")
    out.append({"kind": "semantic_payment", "instruction": (
        "Change every payment period in this agreement (periods within which the Customer must pay) from 30 days to "
        "45 days, keeping each clause's wording style (“thirty (30) days” becomes “forty-five (45) days”, “30 days” "
        "becomes “45 days”). Do not change any other period, such as notice, cure or retention periods, and do not "
        "change the Schedule 2 table."), "gold": gold})

    gold = [re.sub(r"\bClause 9(\.\d+)?\b", lambda m: f"Clause 10{m.group(1) or ''}", t) for t in texts]
    out.append({"kind": "crossref_update", "instruction": (
        "Clause 9 is being renumbered as Clause 10. Update every cross-reference to Clause 9 or its sub-clauses "
        "(e.g. “Clause 9.2” becomes “Clause 10.2”). Do not change the clause headings or numbering themselves, and "
        "do not change references to other clauses such as Clause 19 or 29."), "gold": gold})
    return out
