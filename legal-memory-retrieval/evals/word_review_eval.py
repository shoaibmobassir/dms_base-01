"""Word review on real multi-reviewer files (plan 18): read, resolve, edit, comment — checked
against the files' own XML and LibreOffice.

For every .docx given (default: the reviewed files in ../docs):

  counts_ok        our revisions per author = the raw revision elements in document.xml
  comments_ok      comments / replies / resolved read = the raw comments.xml + commentsExtended
  clear_ok         accepting or rejecting everything leaves no revision (moves included)
  by_person_ok     accepting person by person = accepting everything (same Final text)
  preserve_ok      a browser save on a clean paragraph keeps every other reviewer's revisions
                   exactly (type, author, text) and adds only the editor's
  locked_ok        a save on a paragraph with someone else's pending change is refused
  comment_trip_ok  Precentis comments + a reply to a Word thread + resolving it, written into
                   the file and read back; writing again adds nothing; revisions untouched
  valid_ok         every file we produce opens in python-docx and converts in LibreOffice

    python evals/word_review_eval.py [file.docx ...]
"""
from __future__ import annotations

import io
import json
import sys
import time
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.documents import docx_comments as C  # noqa: E402
from app.documents import docx_review as R  # noqa: E402
from app.documents.editing import EditError, _edit_docx  # noqa: E402
from app.documents.pdf_render import _gotenberg  # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W15 = "http://schemas.microsoft.com/office/word/2012/wordml"
DEFAULT = sorted((ROOT.parent / "docs").glob("*.docx"))
OUT = ROOT / "evals" / "last_word_review.json"
EDITOR = "Eval Editor"


def _raw_counts(data: bytes) -> Counter:
    doc = etree.fromstring(zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml"))
    body = doc.find(f"{{{W}}}body")
    tags = {f"{{{W}}}{t}" for t in ("ins", "del", "moveFrom", "moveTo", "rPrChange", "pPrChange", "tblPrChange", "trPrChange",
                                     "tcPrChange", "tblGridChange", "sectPrChange", "tblPrExChange", "numberingChange",
                                     "cellIns", "cellDel")}
    return Counter(" ".join((el.get(f"{{{W}}}author") or "").split()) or "Unknown" for el in body.iter() if el.tag in tags)


def _raw_comments(data: bytes) -> tuple[int, int, int]:
    z = zipfile.ZipFile(io.BytesIO(data))
    if "word/comments.xml" not in z.namelist():
        return 0, 0, 0
    n = len(etree.fromstring(z.read("word/comments.xml")).findall(f"{{{W}}}comment"))
    replies = done = 0
    if "word/commentsExtended.xml" in z.namelist():
        for e in etree.fromstring(z.read("word/commentsExtended.xml")).iter(f"{{{W15}}}commentEx"):
            replies += e.get(f"{{{W15}}}paraIdParent") is not None
            done += e.get(f"{{{W15}}}done") == "1"
    return n, replies, done


def _valid(data: bytes) -> bool:
    Document(io.BytesIO(data))
    pdf = _gotenberg(data, ".docx")
    return bool(pdf) and pdf[:4] == b"%PDF"


def check(path: Path) -> dict:
    data = path.read_bytes()
    t0 = time.perf_counter()
    revs = R.read_revisions(data)
    read_s = time.perf_counter() - t0
    row: dict = {"file": path.name, "revisions": len(revs), "authors": dict(Counter(r.author for r in revs)),
                 "changes": len(R.group_changes(data, revs)), "read_s": round(read_s, 2)}
    row["counts_ok"] = Counter(r.author for r in revs) == _raw_counts(data)

    comments = R and C.read_comments(data)
    raw_n, raw_replies, raw_done = _raw_comments(data)
    row["comments"] = len(comments)
    row["comments_ok"] = (len(comments), sum(1 for c in comments if c["parent_para_id"]),
                          sum(1 for c in comments if c["done"])) == (raw_n, raw_replies, raw_done)

    accepted, rejected = R.accept_everything(data), R.reject_everything(data)
    row["clear_ok"] = R.read_revisions(accepted) == [] and R.read_revisions(rejected) == []
    cur = data
    for author in sorted({r.author for r in revs}):
        cur, _ = R.resolve(cur, R.keys_by(cur, authors=[author]), accept=True)
    row["by_person_ok"] = R.text_view(cur, "final") == R.text_view(accepted, "final") and R.read_revisions(cur) == []

    # A browser save on the longest clean paragraph, and one on a paragraph with pending changes.
    pending = R.paragraph_pending(data, revs)
    doc = Document(io.BytesIO(data))
    clean = [i for i, p in enumerate(doc.paragraphs) if i not in pending and len(p.text) > 40]
    busy = sorted(pending)
    before = [(r.type, r.author, r.text) for r in revs]
    if clean:
        pid = max(clean, key=lambda i: len(doc.paragraphs[i].text))
        text = doc.paragraphs[pid].text
        t0 = time.perf_counter()
        out, _ = _edit_docx(data, [{"op": "replace", "pid": pid, "text": text + " (reviewed)"}], EDITOR)
        row["save_s"] = round(time.perf_counter() - t0, 2)
        after = R.read_revisions(out)
        row["preserve_ok"] = [(r.type, r.author, r.text) for r in after if r.author != EDITOR] == before \
            and {r.author for r in after if r.author == EDITOR} == {EDITOR} \
            and any(t.endswith(text + " (reviewed)") for t in R.text_view(out, "final"))  # other reviewers' paragraph joins shift positions
        row["valid_save"] = _valid(out)
    else:
        row["preserve_ok"], row["valid_save"] = True, True
    try:
        _edit_docx(data, [{"op": "replace", "pid": busy[0], "text": "x"}], EDITOR) if busy else None
        row["locked_ok"] = not busy
    except EditError as exc:
        row["locked_ok"] = exc.status == 422 and exc.extra.get("locked_pids") == [busy[0]]

    root = next((c for c in comments if not c["parent_para_id"]), None)
    final = R.text_view(data, "final")
    target = next((i for i, t in enumerate(final) if len(t) > 60), 0)
    threads = [{"annotation_id": "CMT-EVAL-1", "parent_id": None, "source": "precentis", "external_id": None,
                "author": "Eval Reviewer", "text": "Check this against the order.", "status": "open",
                "quote": final[target][10:40], "pid": target}]
    if root is not None:
        threads += [{"annotation_id": "CMT-EVAL-W", "parent_id": None, "source": "word", "external_id": C.external_id(root),
                     "author": root["author"], "text": root["text"], "status": "resolved"},
                    {"annotation_id": "CMT-EVAL-2", "parent_id": "CMT-EVAL-W", "source": "precentis", "external_id": None,
                     "author": "Eval Reviewer", "text": "Agreed; resolving."}]
    written_file, written = C.write_comments(data, threads)
    back = C.read_comments(written_file)
    by_para = {c["para_id"]: c for c in back}
    again, more = C.write_comments(written_file, [{**t, "external_id": written.get(t["annotation_id"], t.get("external_id"))}
                                                  for t in threads])
    new_root = by_para.get(written.get("CMT-EVAL-1"), {})
    ok = (len(back) == len(comments) + len(written) and " ".join(new_root.get("quote", "").split()) == " ".join(threads[0]["quote"].split())
          and more == {} and len(C.read_comments(again)) == len(back) and R.read_revisions(written_file) == revs)
    if root is not None:
        reply = by_para.get(written.get("CMT-EVAL-2"), {})
        ok = ok and reply.get("parent_para_id") == root["para_id"] and by_para[root["para_id"]]["done"] is True
    row["comment_trip_ok"] = ok
    row["valid_ok"] = all(_valid(x) for x in (accepted, rejected, cur, written_file)) and row.pop("valid_save")
    row["all_ok"] = all(v for k, v in row.items() if k.endswith("_ok"))
    return row


def main() -> None:
    files = [Path(a) for a in sys.argv[1:]] or DEFAULT
    results = [check(f) for f in files]
    for r in results:
        print(json.dumps({k: v for k, v in r.items() if k != "authors"}, ensure_ascii=False))
        print("   authors:", r["authors"])
    summary = {"files": len(results), "all_ok": all(r["all_ok"] for r in results)}
    OUT.write_text(json.dumps({"summary": summary, "results": results}, indent=1, ensure_ascii=False))
    print(json.dumps(summary))
    sys.exit(0 if summary["all_ok"] else 1)


if __name__ == "__main__":
    main()
