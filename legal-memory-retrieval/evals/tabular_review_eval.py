"""Tabular review accuracy baseline (plan 22, X3) over live HTTP with the real model.

Generates agreements whose key facts are drawn at random (parties, effective date, governing law, notice period,
liability cap, change of control) and written in varied wording next to planted distractors (other day counts, a
second law named for a different purpose, a cap-like amount that is not the cap). Uploads them into a throwaway
project, runs one tabular review with the matching columns, and scores each cell against the gold value.

    python evals/tabular_review_eval.py --base-url http://127.0.0.1:8021 --member MEM-00001 [--docs 20 --seed 11]

CUAD / LegalBench-RAG would be the external benchmark; it needs licence approval before download (plan 14), so this
internal set is the baseline until then. Writes evals/last_tabular_review.json.
"""
from __future__ import annotations

import argparse
import io
import json
import random
import re
import time
from datetime import date
from pathlib import Path

import httpx
from docx import Document

PARTIES = ["Acme Technologies Private Limited", "Norland Shipping AS", "Kestrel Insurance plc", "Zephyrine Robotics Inc.",
           "Obelisk Bank Limited", "Halcyon Foods LLP", "Meridian Power Corporation", "Tessera Analytics Pte. Ltd.",
           "Quorum Health Services", "Vantage Logistics GmbH", "Saffron Textiles Limited", "Polar Data Centres BV"]
LAWS = ["England and Wales", "India", "Singapore", "the State of New York"]
NUM_WORDS = {15: "fifteen", 30: "thirty", 45: "forty-five", 60: "sixty", 90: "ninety"}
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
FILLER = [
    "The Supplier shall perform the Services with reasonable skill, care and diligence in accordance with Good Industry Practice.",
    "Each party shall comply with all applicable laws and regulations in performing its obligations under this Agreement.",
    "The parties shall meet at least quarterly to review performance and agree any improvements required.",
    "Any approval given under this Agreement shall not relieve a party of any of its other obligations.",
    "The Supplier shall maintain complete and accurate records relating to the Services and make them available on request.",
]
COLUMNS = [
    {"label": "Parties", "question": "Who are the parties to this agreement?", "answer_format": "list"},
    {"label": "Effective date", "question": "What is the effective date of this agreement?", "answer_format": "date"},
    {"label": "Governing law", "question": "Which law governs this agreement?", "answer_format": "text"},
    {"label": "Termination notice", "question": "How many days' notice is needed to terminate for convenience?", "answer_format": "number"},
    {"label": "Liability cap", "question": "What is the cap on each party's total liability?", "answer_format": "money"},
    {"label": "Change of control", "question": "Does the agreement give a party rights on a change of control of the other?", "answer_format": "yes_no"},
]


def make(i: int, rng: random.Random) -> tuple[str, bytes, dict]:
    a, b = rng.sample(PARTIES, 2)
    law, other_law = rng.sample(LAWS, 2)
    notice = rng.choice(list(NUM_WORDS))
    cure = rng.choice([d for d in NUM_WORDS if d != notice])
    eff = date(rng.choice([2025, 2026]), rng.randint(1, 12), rng.randint(1, 28))
    cap = rng.choice([50, 75, 100, 250, 500]) * 100_000
    decoy = cap * rng.choice([2, 3])
    coc = rng.random() < 0.5
    d = Document()
    d.add_heading(f"Services Agreement {i + 1}", 0)
    style = rng.random()
    d.add_paragraph(f"This Agreement is dated {eff.day} {MONTHS[eff.month - 1]} {eff.year} (the “Effective Date”) and is made between "
                    f"{a} (the “Customer”) and {b} (the “Supplier”)." if style < 0.5 else
                    f"THIS AGREEMENT is entered into between (1) {a} (“Customer”) and (2) {b} (“Supplier”), with effect from "
                    f"{eff.day} {MONTHS[eff.month - 1]} {eff.year}.")
    sections = []
    sections.append(("Term and termination", [
        f"Either party may terminate this Agreement for convenience by giving not less than {NUM_WORDS[notice]} ({notice}) days' written notice." if rng.random() < 0.5
        else f"This Agreement may be terminated by either party without cause on {notice} days' prior written notice to the other.",
        f"A party may terminate immediately if the other fails to remedy a material breach within {cure} days of notice requiring it to do so.",
    ]))
    sections.append(("Limitation of liability", [
        f"Subject to clause 9.3, each party's total aggregate liability arising under or in connection with this Agreement shall not exceed INR {cap:,}.",
        f"The Supplier shall maintain professional indemnity insurance of not less than INR {decoy:,}.",
    ]))
    if coc:
        sections.append(("Change of control", [
            "If the Supplier undergoes a change of control, the Customer may terminate this Agreement on written notice within ninety days of becoming aware of it."]))
    sections.append(("Data protection", [f"Personal data transferred under this Agreement shall be processed in accordance with the data protection laws of {other_law}."]))
    sections.append(("Governing law", [f"This Agreement and any non-contractual obligations arising out of it are governed by the laws of {law}."]))
    rng.shuffle(sections)
    for n, (title, paras) in enumerate(sections, 1):
        d.add_heading(f"{n}. {title}", 1)
        for _ in range(rng.randint(3, 8)):
            d.add_paragraph(rng.choice(FILLER))
        for p in paras:
            d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    gold = {"Parties": [a, b], "Effective date": eff.isoformat(), "Governing law": law, "Termination notice": notice,
            "Liability cap": cap, "Change of control": coc}
    return f"tr-eval-{i + 1:02d}.docx", buf.getvalue(), gold


def _digits(s: str) -> list[int]:
    return [int(x.replace(",", "")) for x in re.findall(r"\d[\d,]*", s or "")]


def score(label: str, answer: str | None, gold) -> bool:
    a = (answer or "").strip()
    low = a.lower()
    if not a:
        return False
    if label == "Parties":
        return all(g.split()[0].lower() in low for g in gold)
    if label == "Effective date":
        return gold in a or gold.replace("-", "") in re.sub(r"\D", "", a)
    if label == "Governing law":
        key = gold.replace("the State of ", "").lower()
        return key in low
    if label == "Termination notice":
        return gold in _digits(a) or NUM_WORDS[gold] in low
    if label == "Liability cap":
        return gold in _digits(a)
    if label == "Change of control":
        return low.startswith("yes") == gold
    return False


def run(base: str, member: str, n: int, seed: int) -> dict:
    rng = random.Random(seed)
    h = {"X-Member-Id": member}
    c = httpx.Client(base_url=base, headers=h, timeout=120)
    project = c.post("/api/projects", json={"title": f"E2E-TMP tabular eval {seed}"}).json()["project_id"]
    docs = [make(i, rng) for i in range(n)]
    files = [("files", (name, data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")) for name, data, _ in docs]
    batch = c.post("/api/uploads/batches", data={"container_kind": "project", "container_id": project}, files=files).json()
    ran = c.post(f"/api/uploads/batches/{batch['batch_id']}/run").json()
    by_name = {f["relative_path"]: f["document_id"] for f in ran["batch"]["files"]}
    ids = [by_name[name] for name, _, _ in docs]
    t0 = time.perf_counter()
    review = c.post("/api/tabular/reviews", json={"title": "eval", "kind": "project", "id": project, "columns": COLUMNS,
                                                   "document_ids": ids}).json()
    rid = review["review_id"]
    while True:
        view = c.get(f"/api/tabular/reviews/{rid}").json()
        counts = view["counts"]
        if not counts.get("pending") and not counts.get("running"):
            break
        if time.perf_counter() - t0 > 900:
            break
        time.sleep(2)
    elapsed = time.perf_counter() - t0
    cols = {col["column_id"]: col["label"] for col in view["columns"]}
    rows = {r["row_id"]: r["document_id"] for r in view["rows"]}
    gold_by_doc = {doc_id: g for doc_id, (_, _, g) in zip(ids, docs)}
    per = {label: {"ok": 0, "n": 0} for label in cols.values()}
    verified = cited = 0
    misses = []
    for cell in view["cells"]:
        label = cols[cell["column_id"]]
        g = gold_by_doc[rows[cell["row_id"]]][label]
        ok = score(label, cell.get("answer"), g)
        per[label]["n"] += 1
        per[label]["ok"] += ok
        if cell.get("citations"):
            cited += 1
            verified += bool(cell["citations"][0].get("verified"))
        if not ok:
            misses.append({"column": label, "gold": g, "answer": cell.get("answer"), "status": cell["status"]})
    total_ok = sum(v["ok"] for v in per.values())
    total = sum(v["n"] for v in per.values())
    report = {"seed": seed, "documents": n, "cells": total, "accuracy": round(total_ok / total, 3) if total else None,
              "per_column": {k: round(v["ok"] / v["n"], 3) for k, v in per.items()},
              "cited": cited, "verified_quotes": verified, "fill_seconds": round(elapsed, 1),
              "misses": misses[:30], "project": project}
    # clean up the throwaway project (its documents go with it)
    from app.ingest.purge import purge_projects

    purge_projects([project])
    return report


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--base-url", default="http://127.0.0.1:8021")
    ap.add_argument("--member", default="MEM-00001")
    ap.add_argument("--docs", type=int, default=20)
    ap.add_argument("--seed", type=int, default=11)
    args = ap.parse_args()
    report = run(args.base_url, args.member, args.docs, args.seed)
    Path("evals/last_tabular_review.json").write_text(json.dumps(report, indent=2, default=str))
    print(json.dumps({k: v for k, v in report.items() if k != "misses"}, indent=2))
    for m in report["misses"][:12]:
        print("MISS", m)
    return 0


if __name__ == "__main__":
    import sys

    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    raise SystemExit(main())
