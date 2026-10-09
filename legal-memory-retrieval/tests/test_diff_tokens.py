"""Word-level diff: numbers and dates stay whole, punctuation stays out of them, adjacent changes read as one."""
from __future__ import annotations

import io
import re

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.documents.diff import compute_word_redline
from app.documents.editing import _word_diff
from app.documents.tokens import tokenize, word_ops
from app.drafting.docx_tracked import apply_tracked_changes, view


@pytest.mark.parametrize("text", [
    "The Executive shall receive INR 1,20,00,000 and shall be eligible.",
    "with effect from 17 October 2026.",
    "Appeal No. 163 of 2018 (filed 12.02.2025 at 10:30).",
    "Hon’ble Tribunal, issue-wise; 1.25% p.a.; clause 5(a)(ii).",
    "  leading and trailing  ", "", "tabs\tand\nnewlines", "snake_case and ünïcode",
])
def test_tokens_always_join_back_to_the_text(text):
    assert "".join(tokenize(text)) == text


def test_numbers_dates_and_words_are_kept_whole_and_punctuation_is_separate():
    assert tokenize("INR 1,20,00,000.") == ["INR", " ", "1,20,00,000", "."]
    assert tokenize("from 17 October") == ["from", " ", "17", " ", "October"]
    assert tokenize("filed 12.02.2025, at 10:30") == ["filed", " ", "12.02.2025", ",", " ", "at", " ", "10:30"]
    assert tokenize("Hon’ble issue-wise") == ["Hon’ble", " ", "issue-wise"]


def test_changed_amount_is_one_replace_and_the_full_stop_is_untouched():
    a, b = tokenize("compensation of INR 1,20,00,000."), tokenize("compensation of INR 1,50,00,000.")
    changes = [(tag, "".join(a[i1:i2]), "".join(b[j1:j2])) for tag, i1, i2, j1, j2 in word_ops(a, b) if tag != "equal"]
    assert changes == [("replace", "1,20,00,000", "1,50,00,000")]


def test_changed_day_is_a_whole_number_never_a_digit():
    a, b = tokenize("with effect from 1 October 2026"), tokenize("with effect from 7 October 2026")
    changes = [(tag, "".join(a[i1:i2]), "".join(b[j1:j2])) for tag, i1, i2, j1, j2 in word_ops(a, b) if tag != "equal"]
    assert changes == [("replace", "1", "7")]
    a, b = tokenize("effect from 17 October"), tokenize("effect from 7 October")
    assert [(t, "".join(a[i1:i2]), "".join(b[j1:j2])) for t, i1, i2, j1, j2 in word_ops(a, b) if t != "equal"] == [("replace", "17", "7")]


def test_changes_separated_only_by_a_space_are_one_change():
    a, b = tokenize("the parties shall pay"), tokenize("each party must pay")
    changes = [(tag, "".join(a[i1:i2]), "".join(b[j1:j2])) for tag, i1, i2, j1, j2 in word_ops(a, b) if tag != "equal"]
    assert changes == [("replace", "the parties shall", "each party must")]


def test_two_separate_deletions_are_not_merged_into_a_replaced_space():
    a, b = tokenize("one red big house"), tokenize("one house")
    kinds = [tag for tag, *_ in word_ops(a, b) if tag != "equal"]
    assert kinds == ["delete"]            # "red big " is one contiguous deletion
    a, b = tokenize("alpha X beta Y gamma"), tokenize("alpha beta gamma")
    assert [tag for tag, *_ in word_ops(a, b) if tag != "equal"] == ["delete", "delete"]


@pytest.mark.parametrize("old,new", [
    ("The fee is INR 1,20,00,000.", "The fee is INR 1,50,00,000."),
    ("from 17 October 2026", "from 7 October 2026"),
    ("the parties shall pay", "each party must pay"),
    ("one red big house", "one house"),
    ("a b c d", "a x c y"),
    ("", "new text"), ("old text", ""),
])
def test_applying_the_ops_reproduces_the_new_text_and_reject_the_old(old, new):
    a, b = tokenize(old), tokenize(new)
    rebuilt = "".join(
        "".join(a[i1:i2]) if tag == "equal" else "".join(b[j1:j2]) for tag, i1, i2, j1, j2 in word_ops(a, b))
    assert rebuilt == new
    rejected = "".join("".join(a[i1:i2]) for tag, i1, i2, j1, j2 in word_ops(a, b))
    assert rejected == old


def test_compare_segments_and_redline_use_the_same_tokens():
    seg = _word_diff("fee INR 1,20,00,000.", "fee INR 1,50,00,000.")
    assert [(s["t"], s["text"]) for s in seg] == [("eq", "fee INR "), ("del", "1,20,00,000"), ("ins", "1,50,00,000"), ("eq", ".")]
    red = compute_word_redline("fee INR 1,20,00,000.", "fee INR 1,50,00,000.")
    assert [(t.token_type, t.text) for t in red if t.token_type != "equal"] == [("delete", "1,20,00,000"), ("insert", "1,50,00,000")]


def _docx(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_tracked_change_marks_only_the_number_and_accepts_and_rejects_cleanly():
    old = "The Executive shall receive fixed annual compensation of INR 1,20,00,000."
    new = old.replace("1,20,00,000", "1,50,00,000")
    data, stats = apply_tracked_changes(_docx(old), [{"op": "replace", "pid": 0, "text": new}])
    assert stats["replaced"] == 1
    assert view(data, accept=True)[0] == new
    assert view(data, accept=False)[0] == old
    body = Document(io.BytesIO(data)).element.body
    deleted = ["".join(t.text or "" for t in d.iter(qn("w:delText"))) for d in body.iter(qn("w:del"))]
    inserted = ["".join(t.text or "" for t in i.iter(qn("w:t"))) for i in body.iter(qn("w:ins"))]
    assert deleted == ["1,20,00,000"] and inserted == ["1,50,00,000"]    # not the full stop, not the sentence


def test_a_phrase_change_is_one_deletion_and_one_insertion():
    old, new = "The parties shall pay within thirty days.", "Each party must pay within thirty days."
    data, _ = apply_tracked_changes(_docx(old), [{"op": "replace", "pid": 0, "text": new}])
    body = Document(io.BytesIO(data)).element.body
    assert len(list(body.iter(qn("w:del")))) == 1 and len(list(body.iter(qn("w:ins")))) == 1
    assert view(data, accept=True)[0] == new and view(data, accept=False)[0] == old
