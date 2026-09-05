"""Tests for PDF/DOCX → Page → Section → Block and version-scoped hierarchical chunks."""
from __future__ import annotations

import os
from pathlib import Path

import pytest

from app.documents.canonical import parse_canonical_blocks, parse_from_extracted
from app.documents.hierarchical_chunks import (
    HierarchicalChunk,
    build_hierarchical_chunks,
)
from app.ingest.extractors.types import PageSpan, join_pages


class TestPageAwareBlocks:
    def test_page_spans_drive_page_numbers(self):
        # Two pages joined with form-feed
        p1 = "ARTICLE I DEFINITIONS\n\nBuyer means the purchasing party."
        p2 = "Section 8.2 Tax Indemnity. Cap shall not exceed $75,000."
        extracted = join_pages([p1, p2])
        assert extracted.page_count == 2
        assert "\f" in extracted.text

        blocks = parse_from_extracted(extracted, "DOC-T", "VER-T")
        assert blocks
        # First blocks on page 1, indemnity content on page 2
        pages = {b.page_number for b in blocks}
        assert 1 in pages and 2 in pages
        assert any(b.metadata.get("page_source") == "span" for b in blocks)
        assert any("$75,000" in b.text for b in blocks)

    def test_md_section_heading(self):
        body = "## Section 8.2. Tax Indemnity\n\nThe Seller shall indemnify Buyer."
        blocks = parse_canonical_blocks(body, "DOC-1", "VER-1")
        assert any(b.block_type == "heading" and b.section_id == "8.2" for b in blocks)


class TestHierarchicalChunks:
    def test_parent_child_and_envelope(self):
        body = (
            "ARTICLE VIII INDEMNIFICATION\n\n"
            "Section 8.1 Survival. Reps survive 24 months.\n\n"
            "Section 8.2 Tax Indemnity. Cap is $50,000.\n\n"
            "More indemnity language that continues for the child chunk sizing."
        )
        blocks = parse_canonical_blocks(body, "DOC-H", "VER-H")
        chunks = build_hierarchical_chunks(
            blocks,
            document_id="DOC-H",
            version_id="VER-H",
            matter_id="MTR-H",
            folder_path="Client/Matter/Agreements",
        )
        parents = [c for c in chunks if c.is_parent]
        children = [c for c in chunks if not c.is_parent]
        assert parents
        assert children
        assert all(c.version_id == "VER-H" for c in chunks)
        assert all(c.folder_path.endswith("Agreements") for c in chunks)
        env = children[0].context_envelope(
            client_name="Acme",
            matter_name="Project Ganges",
            document_name="SPA.pdf",
            version_number=3,
        )
        assert "CLIENT: Acme" in env
        assert "FOLDER: Client/Matter/Agreements" in env
        assert "VERSION: v3" in env
        assert "SECTION:" in env


class TestDocxExtractor:
    def test_docx_roundtrip(self, tmp_path: Path):
        pytest.importorskip("docx")
        from docx import Document

        from app.ingest.extractors.docx import DocxExtractor

        path = tmp_path / "sample.docx"
        doc = Document()
        doc.add_heading("SECTION 1 Sale and Purchase", level=1)
        doc.add_paragraph("The Seller agrees to sell the Shares to the Buyer.")
        doc.add_paragraph("Section 8.2 Liability. Aggregate liability shall not exceed $100,000.")
        doc.save(path)

        extracted = DocxExtractor().extract_structured(str(path))
        assert extracted.source_format == "docx"
        assert extracted.page_count >= 1
        assert "Shares" in extracted.text
        blocks = parse_from_extracted(extracted, "DOC-X", "VER-X")
        assert any(b.block_type in ("heading", "clause", "paragraph") for b in blocks)


class TestUploadCreatesVersionChunks:
    def test_upload_writes_version_scoped_chunks(self, tmp_path: Path, monkeypatch):
        from app.config import settings
        from app.db.connection import connect
        from app.documents.hierarchical_chunks import get_version_chunks
        from app.ingest.upload_batch import create_upload_batch, process_upload_batch
        from app.storage.object_store import reset_object_store_for_tests

        reset_object_store_for_tests()
        monkeypatch.setattr(settings, "object_store_backend", "local")
        monkeypatch.setattr(settings, "object_store_root", str(tmp_path / "obj"))

        with connect() as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT matter_id FROM matters LIMIT 1")
                row = cur.fetchone()
        if not row:
            pytest.skip("no matters")
        matter_id = row["matter_id"] if isinstance(row, dict) else row[0]

        uniq = os.urandom(3).hex()
        body = (
            f"ARTICLE VIII INDEMNIFICATION {uniq}\n\n"
            "Section 8.2 Tax Indemnity. Cap shall not exceed $75,000.\n\n"
            "IN WITNESS WHEREOF the parties execute this Agreement."
        )
        batch = create_upload_batch(
            matter_id=matter_id,
            files=[(f"Agreements/SPA_{uniq}.txt", body.encode())],
        )
        result = process_upload_batch(batch["batch_id"])
        assert result["indexed"] == 1

        with connect() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    SELECT current_version_id FROM documents
                    WHERE title = %(t)s ORDER BY ingested_at DESC NULLS LAST LIMIT 1
                    """,
                    {"t": f"SPA_{uniq}.txt"},
                )
                doc = cur.fetchone()
        vid = doc["current_version_id"] if isinstance(doc, dict) else doc[0]
        chunks = get_version_chunks(vid)
        assert chunks
        assert all(c["version_id"] == vid for c in chunks)
        assert any(c.get("section_id") for c in chunks)
        children = get_version_chunks(vid, children_only=True)
        assert children
