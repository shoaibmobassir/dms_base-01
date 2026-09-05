"""
Comprehensive Test Suite for Integrated Mike Features in FirmOS / DMS Knowledge Base.
Clean-room verification covering:
1. Key Vault & AES-256-GCM Encryption
2. Tabular Reviews & Excel Export
3. Legal Workflows & YAML DAG Execution
4. Word DOCX Native Tracked Changes (<w:ins>, <w:del>) & Clause Diff
5. Word Taskpane Auth Handoff & Selection Endpoints
6. Case Law Citation Tokenizer & CourtListener Client
7. Tamper-Evident Manifest Signing & ZIP Packaging
"""

import io
import json
import zipfile
import pytest
from fastapi.testclient import TestClient
import openpyxl

from app.api.main import app
from app.audit.manifest_signer import get_manifest_signer
from app.auth.handoff import get_auth_handoff_service
from app.auth.key_vault import KeyVault
from app.caselaw.citation_parser import get_citation_parser
from app.caselaw.courtlistener_client import get_courtlistener_client
from app.drafting.clause_diff_service import get_clause_diff_service
from app.drafting.docx_redline_generator import DocxRedlineGenerator
from app.review.spreadsheet_exporter import SpreadsheetExporter
from app.workflows.catalog_loader import get_catalog_loader
from app.workflows.engine import get_workflow_engine


@pytest.fixture
def client():
    return TestClient(app)


# -----------------------------------------------------------------------------
# 1. Key Vault (AES-256-GCM)
# -----------------------------------------------------------------------------
def test_key_vault_encryption_roundtrip():
    vault = KeyVault(master_secret="test-secret-key-vault-firmos-32b")
    raw_key = "sk-ant-api03-abcdef1234567890-secure"
    encrypted = vault.encrypt_key(raw_key)
    assert encrypted != raw_key
    decrypted = vault.decrypt_key(encrypted)
    assert decrypted == raw_key


def test_key_vault_tamper_detection():
    vault = KeyVault(master_secret="test-secret-key-vault-firmos-32b")
    raw_key = "sk-openai-test-key-999"
    encrypted = vault.encrypt_key(raw_key)
    tampered = encrypted[:-4] + "AAAA"
    with pytest.raises(ValueError):
        vault.decrypt_key(tampered)


# -----------------------------------------------------------------------------
# 2. Tabular Reviews & Excel Export
# -----------------------------------------------------------------------------
def test_spreadsheet_exporter():
    columns = [
        {"id": "c1", "label": "Governing Law"},
        {"id": "c2", "label": "Liability Cap"},
    ]
    rows = [{"document_id": "DOC-001", "title": "Master NDA"}]
    cells = {
        "DOC-001": {
            "c1": {"value": "English Law", "confidence": 0.95, "citations": ["CHK-01"], "reasoning": "Clause 14.1"},
            "c2": {"value": "£1,000,000", "confidence": 0.90, "citations": ["CHK-02"], "reasoning": "Clause 9.2"},
        }
    }
    xlsx_bytes = SpreadsheetExporter.export_xlsx("Test DD Review", columns, rows, cells)
    assert len(xlsx_bytes) > 0

    # Verify openpyxl can load and read sheets
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "Review Matrix" in wb.sheetnames
    assert "Evidence & Citations" in wb.sheetnames
    ws = wb["Review Matrix"]
    assert ws["A1"].value.startswith("FirmOS Legal Intelligence")
    assert ws["C4"].value == "English Law"


def test_tabular_review_api_workflow(client):
    create_payload = {
        "title": "Vendor Contract Audit",
        "document_ids": ["DOC-00001", "DOC-00002"],
        "columns": [
            {"id": "col_law", "label": "Governing Law", "prompt": "Identify governing law"},
            {"id": "col_cap", "label": "Liability Cap", "prompt": "Extract liability cap", "data_type": "currency"},
        ],
    }
    resp = client.post("/api/tabular/reviews", json=create_payload)
    assert resp.status_code == 200
    data = resp.json()
    review_id = data["review_id"]
    assert data["title"] == "Vendor Contract Audit"
    assert "DOC-00001" in data["cells"]

    # Test override
    patch_resp = client.patch(
        f"/api/tabular/reviews/{review_id}/cells/DOC-00001/col_law",
        json={"value": "New York Law", "reasoning": "Senior Partner review"},
    )
    assert patch_resp.status_code == 200
    assert patch_resp.json()["value"] == "New York Law"
    assert patch_resp.json()["is_overridden"] is True

    # Test Excel download
    export_resp = client.get(f"/api/tabular/reviews/{review_id}/export/xlsx")
    assert export_resp.status_code == 200
    assert export_resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# -----------------------------------------------------------------------------
# 3. Legal Playbooks & Workflow Engine
# -----------------------------------------------------------------------------
def test_workflow_catalog_loader():
    loader = get_catalog_loader()
    wfs = loader.list_workflows()
    assert len(wfs) >= 3
    ids = [w.id for w in wfs]
    assert "cease-and-desist-drafter" in ids
    assert "contract-triage-router" in ids
    assert "master-services-agreement-review" in ids


def test_workflow_api_execution(client):
    req_payload = {
        "workflow_id": "cease-and-desist-drafter",
        "inputs": {
            "rights_holder": "Acme Software Corp",
            "infringing_party": "Pirate Tech Ltd",
            "infringing_details": "Unauthorized reproduction of proprietary source code",
            "remedy_deadline_days": 7,
        },
    }
    resp = client.post("/api/workflows/run", json=req_payload)
    assert resp.status_code == 200
    run_data = resp.json()
    assert run_data["status"] == "completed"
    assert "Acme Software Corp" in run_data["final_output"] or "LEGAL DEMAND" in run_data["final_output"]
    assert "retrieve_firm_cd_precedents" in run_data["step_outputs"]


# -----------------------------------------------------------------------------
# 4. DOCX Track Changes & Clause Diff
# -----------------------------------------------------------------------------
def test_docx_tracked_changes_generator():
    gen = DocxRedlineGenerator(author="FirmOS AI")
    orig = "Customer shall pay within 30 days of invoice receipt."
    rev = "Customer shall pay within 45 business days of verified invoice receipt."
    docx_bytes = gen.create_tracked_diff_docx(orig, rev, title="Payment Terms Redline")
    assert len(docx_bytes) > 0
    # Must be a valid zip / docx archive
    zf = zipfile.ZipFile(io.BytesIO(docx_bytes))
    assert "word/document.xml" in zf.namelist()
    doc_xml = zf.read("word/document.xml").decode("utf-8")
    assert "w:ins" in doc_xml
    assert "w:del" in doc_xml


@pytest.mark.asyncio
async def test_clause_diff_service():
    diff_svc = get_clause_diff_service()
    clause = "Supplier liability shall be uncapped for all indirect and consequential losses."
    res = await diff_svc.analyze_and_redline(
        clause_text=clause,
        clause_type="Limitation of Liability",
        client_stance="Supplier",
    )
    assert res.original_clause == clause
    assert res.risk_severity in ("low", "medium", "high", "critical")
    assert len(res.proposed_redline) > 0


# -----------------------------------------------------------------------------
# 5. Word Taskpane Auth Handoff & Selection Router
# -----------------------------------------------------------------------------
def test_auth_handoff_ticket_lifecycle():
    svc = get_auth_handoff_service()
    ticket = svc.create_ticket("MEM-00001", "lawyer@firm.com")
    assert ticket.startswith("TKT-")

    # Exchange once -> succeeds
    session = svc.exchange_ticket(ticket)
    assert session is not None
    assert session["member_id"] == "MEM-00001"

    # Replay -> fails (one-time token)
    assert svc.exchange_ticket(ticket) is None


def test_word_selection_api(client):
    resp = client.post(
        "/api/word/analyze-selection",
        json={"selection_text": "Governing law shall be English Law.", "action": "explain"},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["action"] == "explain"
    assert "analysis" in data


# -----------------------------------------------------------------------------
# 6. Case Law & Citation Verification
# -----------------------------------------------------------------------------
def test_citation_parser():
    parser = get_citation_parser()
    text = (
        "Under Miranda v. Arizona, 384 U.S. 436 (1966) and 123 F.3d 456, "
        "as well as 42 U.S.C. § 1983, constitutional rights are enforceable."
    )
    citations = parser.extract_citations(text)
    assert len(citations) >= 3
    normalized = [c.normalized_citation for c in citations]
    assert any("384 U.S. 436" in n for n in normalized)
    assert any("123 F.3d 456" in n for n in normalized)
    assert any("42 U.S.C. § 1983" in n for n in normalized)


@pytest.mark.asyncio
async def test_courtlistener_client():
    client = get_courtlistener_client()
    opinion = await client.verify_citation("384 U.S. 436")
    assert opinion is not None
    assert opinion.is_good_law is True
    assert opinion.precedential_status in ("Precedential", "Published")



# -----------------------------------------------------------------------------
# 7. Tamper-Evident Export Packaging & Cryptographic Manifest Verification
# -----------------------------------------------------------------------------
def test_manifest_signer_and_verifier():
    signer = get_manifest_signer()
    files = [
        ("report.md", b"# Legal Due Diligence Report\nAll clear."),
        ("data.csv", b"doc_id,risk\nDOC-1,low\nDOC-2,high\n"),
    ]
    meta = {"matter_id": "MAT-00123", "lead_lawyer": "Partner Jane"}

    zip_bytes = signer.build_signed_export_zip("Due Diligence Export", files, meta)
    assert len(zip_bytes) > 0

    zf = zipfile.ZipFile(io.BytesIO(zip_bytes))
    assert "MANIFEST.json" in zf.namelist()
    assert "report.md" in zf.namelist()
    assert "data.csv" in zf.namelist()

    manifest_json = zf.read("MANIFEST.json").decode("utf-8")
    assert signer.verify_manifest(manifest_json) is True

    # Test tampering
    tampered_manifest = manifest_json.replace("Due Diligence Export", "Forged Document")
    assert signer.verify_manifest(tampered_manifest) is False
