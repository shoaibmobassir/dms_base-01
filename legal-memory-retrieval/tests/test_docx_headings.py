"""Word headings: the indexed text is the document's own words, and headings are still found (2026-10-04).

The DOCX extractor used to write every Heading-style paragraph as "SECTION <text>". The reader, search and the Assistant
then showed and quoted "SECTION 5. Remuneration", which is not in the file, so an edit to it could never be placed.
"""
from __future__ import annotations

import io

from docx import Document

from app.chat.doc_nav import outline
from app.documents.canonical import caps_heading, numbered_heading, parse_canonical_blocks
from app.ingest.extractors.docx import DocxExtractor, docx_heading_texts
from tests.test_document_editor import doc, people  # noqa: F401  (fixtures)


def _cto() -> bytes:
    doc = Document()
    doc.add_heading("Employment Agreement - CTO", level=0)
    doc.add_heading("EMPLOYMENT AGREEMENT — CHIEF TECHNOLOGY OFFICER", level=1)
    doc.add_heading("1. Parties", level=2)
    doc.add_paragraph('This Agreement is made between Acme Technologies Private Limited (the "Company") and Mr Arvind Rao.')
    doc.add_heading("5. Remuneration", level=2)
    doc.add_paragraph("The Executive shall receive fixed annual compensation of INR 2,50,00,000.")
    doc.add_heading("Specific disclosure", level=2)
    doc.add_paragraph("Nothing further is disclosed.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_the_extracted_text_is_the_documents_own_words():
    out = DocxExtractor().extract_bytes(_cto())
    assert "SECTION" not in out.text
    assert "5. Remuneration" in out.text and "EMPLOYMENT AGREEMENT — CHIEF TECHNOLOGY OFFICER" in out.text
    assert out.headings == ["Employment Agreement - CTO", "EMPLOYMENT AGREEMENT — CHIEF TECHNOLOGY OFFICER",
                            "1. Parties", "5. Remuneration", "Specific disclosure"]
    assert docx_heading_texts(_cto()) == out.headings


def test_styled_headings_become_heading_blocks_with_their_own_text():
    out = DocxExtractor().extract_bytes(_cto())
    blocks = parse_canonical_blocks(out.text, "DOC-T", "VER-T", page_spans=out.pages, headings=out.headings)
    heads = {b.text: b for b in blocks if b.block_type == "heading"}
    assert set(heads) == set(out.headings)
    assert heads["5. Remuneration"].section_id == "5" and heads["5. Remuneration"].section_title == "Remuneration"
    assert heads["Specific disclosure"].section_id == "specific-disclosure"
    body = next(b for b in blocks if b.text.startswith("The Executive shall"))
    assert body.block_type == "paragraph" and body.section_id == "5"


def test_text_alone_still_finds_numbered_and_capitals_headings():
    text = "\n\n".join(["EMPLOYMENT AGREEMENT", "5. Remuneration", "The Executive shall be paid monthly.",
                        "5. The Executive shall be paid monthly.", "2.1 Definitions"])
    kinds = {b.text: b.block_type for b in parse_canonical_blocks(text, "DOC-T", "VER-T")}
    assert kinds["EMPLOYMENT AGREEMENT"] == "heading" and kinds["5. Remuneration"] == "heading"
    assert kinds["2.1 Definitions"] == "heading"
    assert kinds["5. The Executive shall be paid monthly."] != "heading"  # a numbered sentence is not a heading


def test_rules_for_headings_in_plain_text():
    assert numbered_heading("5. Remuneration") == ("5", "Remuneration")
    assert numbered_heading("12.3 Change of Control") == ("12.3", "Change of Control")
    assert numbered_heading("5. The Executive shall be paid on the last day of each month.") is None
    assert caps_heading("EMPLOYMENT AGREEMENT — CHIEF TECHNOLOGY OFFICER")
    assert not caps_heading("IN WITNESS WHEREOF the parties have signed.") and not caps_heading("NDA")


def test_text_indexed_with_the_old_label_still_parses_the_same():
    blocks = parse_canonical_blocks("SECTION Specific disclosure\n\nNothing further.", "DOC-T", "VER-T")
    assert blocks[0].block_type == "heading" and blocks[0].section_title == "Specific disclosure"


def test_the_assistant_outline_finds_the_sections_of_the_new_text():
    text = "[Page 1]\n" + DocxExtractor().extract_bytes(_cto()).text
    titles = [s.title for s in outline(text)]
    assert "5. Remuneration" in titles and "1. Parties" in titles
    assert "EMPLOYMENT AGREEMENT — CHIEF TECHNOLOGY OFFICER" in titles


def test_reindex_rereads_text_stored_with_the_old_label(doc, people):  # noqa: F811
    from app.db.connection import connect
    from app.documents.docx_reindex import affected_documents, apply, plan

    with connect() as conn:
        conn.execute("UPDATE document_versions SET body = replace(body, 'Article 2 — Fees', 'SECTION Article 2 — Fees') "
                     "WHERE document_id = %s", (doc,))
        conn.commit()
        found = affected_documents(conn, [doc])
        assert [d["document_id"] for d in found] == [doc]
        p = plan(conn, found[0])
        assert len(p.versions) == 1 and "SECTION" not in p.versions[0]["after"]
        out = apply(conn, p)
        assert out["versions"] == 1 and out["chunks"] > 0
        assert affected_documents(conn, [doc]) == []
        heads = [r["text"] for r in conn.execute(
            "SELECT b.text FROM document_blocks b JOIN documents d ON d.current_version_id = b.version_id "
            "WHERE b.document_id = %s AND b.block_type = 'heading' ORDER BY b.sequence", (doc,))]
    assert "Article 2 — Fees" in heads and not any(h.startswith("SECTION") for h in heads)
