"""Tracked formatting changes written into Word files (plan 17, G1) — unit level, no database."""
from __future__ import annotations

import io

from docx import Document
from docx.oxml.ns import qn

from app.documents.docx_format import accept_formatting, apply_formatting, reject_formatting
from app.drafting.docx_tracked import accept_all, apply_tracked_changes, formatting_signature, view

A = "The Client shall pay the Fees within forty-five (45) days."
B = "Either Party may terminate on notice."


def _docx() -> bytes:
    d = Document()
    d.add_paragraph("Services Agreement", style="Title")
    p = d.add_paragraph()
    p.add_run("The Client shall pay the ")
    p.add_run("Fees").bold = True
    p.add_run(" within forty-five (45) days.")
    d.add_paragraph(B)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xml(data: bytes) -> str:
    import zipfile

    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.read("word/document.xml").decode()


def _paras(data: bytes):
    return Document(io.BytesIO(data)).paragraphs


def _flags(p) -> list[tuple[str, bool, bool, bool]]:
    return [(r.text, bool(r.bold), bool(r.italic), bool(r.underline)) for r in p.runs if r.text]


def _bold_at(p, word: str) -> bool:
    for r in p.runs:
        if word in r.text:
            return bool(r.bold)
    raise AssertionError(word)


def test_bold_a_phrase_is_a_tracked_formatting_change():
    src = _docx()
    runs = [{"text": "The Client shall pay the "}, {"text": "Fees", "bold": True}, {"text": " within "},
            {"text": "forty-five (45) days", "bold": True}, {"text": "."}]
    out, stats = apply_formatting(src, [{"op": "format", "pid": 1, "runs": runs}], author="Amina Rahman")
    assert stats["formatted_runs"] >= 1 and stats["format_skipped"] == 0
    xml = _xml(out)
    assert "w:rPrChange" in xml and 'w:author="Amina Rahman"' in xml
    p = _paras(accept_formatting(out))[1]
    assert p.text == A
    assert _bold_at(p, "forty-five") and _bold_at(p, "Fees") and not _bold_at(p, "within")
    # Rejecting restores the original runs exactly.
    old = _paras(src)[1]._p
    back = _paras(reject_formatting(out))[1]._p
    assert formatting_signature(back) == formatting_signature(old)


def test_unbold_and_italic_and_underline():
    src = _docx()
    runs = [{"text": "The Client shall pay the ", "underline": True}, {"text": "Fees", "italic": True},
            {"text": " within forty-five (45) days."}]
    out, _ = apply_formatting(src, [{"op": "format", "pid": 1, "runs": runs}], author="X")
    p = _paras(accept_formatting(out))[1]
    assert not _bold_at(p, "Fees")
    fees = next(r for r in p.runs if r.text == "Fees")
    assert fees.italic and not fees.bold
    assert next(r for r in p.runs if "Client" in r.text).underline


def test_style_change_is_tracked_and_reversible():
    src = _docx()
    out, stats = apply_formatting(src, [{"op": "format", "pid": 2, "style": "Heading 1"}], author="X")
    assert stats["restyled"] == 1 and "w:pPrChange" in _xml(out)
    assert _paras(accept_formatting(out))[2].style.name == "Heading 1"
    assert _paras(reject_formatting(out))[2].style.name == "Normal"


def test_text_and_formatting_in_one_paragraph():
    """Text diff first (docx_tracked), then formatting over the accepted text: the inserted
    words are bold as part of the insertion, unchanged words keep their formatting."""
    src = _docx()
    new = "The Client shall pay the Fees within thirty (30) days."
    ops = [{"op": "replace", "pid": 1, "text": new,
            "runs": [{"text": "The Client shall pay the "}, {"text": "Fees", "bold": True}, {"text": " within "},
                     {"text": "thirty (30)", "bold": True}, {"text": " days."}]}]
    tracked, _ = apply_tracked_changes(src, ops, author="X")
    out, stats = apply_formatting(tracked, ops, author="X")
    assert stats["format_skipped"] == 0
    acc = _paras(accept_formatting(accept_all(out)))[1]
    assert acc.text == new and _bold_at(acc, "thirty") and _bold_at(acc, "Fees") and not _bold_at(acc, "within")
    assert view(out, accept=False)[1] == A  # the text change still rejects cleanly


def test_inserted_paragraph_gets_style_and_marks_without_formatting_revisions():
    src = _docx()
    ops = [{"op": "insert_after", "pid": 2, "text": "Governing law", "style": "Heading 2",
            "runs": [{"text": "Governing law", "bold": True}]},
           {"op": "insert_after", "pid": 2, "text": "English law applies.", "runs": [{"text": "English law applies."}]}]
    tracked, _ = apply_tracked_changes(src, ops, author="X")
    out, stats = apply_formatting(tracked, ops, author="X")
    ps = _paras(accept_formatting(accept_all(out)))
    assert [p.text for p in ps[3:]] == ["Governing law", "English law applies."]
    assert ps[3].style.name == "Heading 2" and _bold_at(ps[3], "Governing")
    # New text is an insertion: no separate formatting revision is recorded for it.
    new_p = Document(io.BytesIO(out)).paragraphs[3]._p
    assert new_p.find(f".//{qn('w:rPrChange')}") is None and new_p.find(f".//{qn('w:pPrChange')}") is None


def test_mismatched_runs_are_skipped_not_guessed():
    out, stats = apply_formatting(_docx(), [{"op": "format", "pid": 1, "runs": [{"text": "Something else", "bold": True}]}],
                                  author="X")
    assert stats["format_skipped"] == 1 and stats["formatted_runs"] == 0
    assert "w:rPrChange" not in _xml(out)


def test_untouched_paragraphs_keep_their_formatting():
    src = _docx()
    out, _ = apply_formatting(src, [{"op": "format", "pid": 2, "style": "Heading 1"}], author="X")
    for i in (0, 1):
        assert formatting_signature(_paras(out)[i]._p) == formatting_signature(_paras(src)[i]._p)
