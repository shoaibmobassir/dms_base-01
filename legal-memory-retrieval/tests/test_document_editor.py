"""Document editor, file versions, locks and compare through the HTTP API (plan 16, E4).

Each test works on a throwaway Word document created in a real open matter and deleted
afterwards; people and the matter come from the database.
"""
from __future__ import annotations

import io
import uuid

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.db.connection import connect
from app.drafting.docx_tracked import formatting_signature, view
from tests.conftest import as_member

PARAS = [
    ("Title", "Services Agreement", []),
    ("Heading 1", "Article 1 — Term", []),
    ("Normal", "This Agreement starts on the Effective Date and runs for two (2) years.", ["Effective Date"]),
    ("Normal", "Either Party may terminate on thirty (30) days' written notice.", ["Party"]),
    ("Heading 1", "Article 2 — Fees", []),
    ("Normal", "The Client shall pay the Fees within forty-five (45) days of invoice.", ["Client", "Fees"]),
    ("Normal", "Late payments bear interest at two per cent (2%) per month.", []),
]


def _docx() -> bytes:
    doc = Document()
    for style, text, bold in PARAS:
        if style == "Title":
            doc.add_heading(text, level=0)
        elif style.startswith("Heading"):
            doc.add_heading(text, level=1)
        else:
            p = doc.add_paragraph()
            rest = text
            for term in bold:
                before, _, rest = rest.partition(term)
                p.add_run(before)
                p.add_run(term).bold = True
            p.add_run(rest)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="module")
def people(seeded):
    with connect() as conn:
        admins = {r["member_id"] for r in conn.execute(
            "SELECT DISTINCT member_id FROM member_roles WHERE role_key IN ('firm_admin', 'risk_compliance')")}
        for m in conn.execute(
            """SELECT m.matter_id, m.client_id FROM matters m JOIN matter_access a USING (matter_id)
               WHERE a.mode = 'open' ORDER BY m.matter_id""").fetchall():
            team = [r["member_id"] for r in conn.execute(
                "SELECT member_id FROM matter_members WHERE matter_id = %s ORDER BY member_id", (m["matter_id"],))]
            editors = [t for t in team if t not in admins]
            outsiders = [r["member_id"] for r in conn.execute(
                "SELECT member_id FROM members WHERE NOT (member_id = ANY(%s)) ORDER BY member_id", (team,))
                         if r["member_id"] not in admins]
            if len(editors) >= 2 and outsiders:
                return {"matter_id": m["matter_id"], "client_id": m["client_id"], "editor": editors[0],
                        "editor2": editors[1], "outsider": outsiders[0]}
    pytest.skip("no open matter with two non-admin staff and an outsider")


@pytest.fixture
def doc(people):
    """A fresh Word document (v1) in the matter; removed after the test."""
    from app.documents import create_version
    from app.storage.object_store import get_object_store

    doc_id = f"DOC-{uuid.uuid4().hex[:10].upper()}"
    data = _docx()
    uri = get_object_store().put(f"tests/{doc_id}/v1.docx", data, content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    with connect() as conn:
        conn.execute(
            """INSERT INTO documents (document_id, matter_id, client_id, title, document_type, body, mime_type, source_uri)
               VALUES (%s, %s, %s, 'Services Agreement.docx', 'Agreement', '', %s, %s)""",
            (doc_id, people["matter_id"], people["client_id"],
             "application/vnd.openxmlformats-officedocument.wordprocessingml.document", uri),
        )
        conn.commit()
    create_version(document_id=doc_id, body="\n".join(t for _, t, _ in PARAS), author_name="Test", source="upload",
                   version_status="draft", storage_uri=uri)
    yield doc_id
    from app.documents.editing import wait_for_indexing

    wait_for_indexing()  # background embeddings finish before the rows go
    with connect() as conn:
        for table in ("document_events", "document_drafts", "document_locks", "chunks", "document_blocks", "version_diffs"):
            conn.execute(f"DELETE FROM {table} WHERE document_id = %s", (doc_id,))
        conn.execute("UPDATE documents SET current_version_id = NULL WHERE document_id = %s", (doc_id,))
        conn.execute("DELETE FROM document_versions WHERE document_id = %s", (doc_id,))
        conn.execute("DELETE FROM documents WHERE document_id = %s", (doc_id,))
        conn.commit()


def _model(client, doc_id, member):
    r = client.get(f"/api/editor/documents/{doc_id}", headers=as_member(member))
    assert r.status_code == 200, r.text
    return r.json()


def _file(doc_id: str) -> bytes:
    from app.storage.object_store import get_object_store

    with connect() as conn:
        uri = conn.execute(
            "SELECT v.storage_uri FROM documents d JOIN document_versions v ON v.version_id = d.current_version_id WHERE d.document_id = %s",
            (doc_id,)).fetchone()["storage_uri"]
    return get_object_store().get(uri)


def _save(client, doc_id, member, base, ops, **kw):
    return client.post(f"/api/editor/documents/{doc_id}/save", headers=as_member(member),
                       json={"base_version_id": base, "ops": ops, **kw})


# ── edit model and permissions ───────────────────────────────────────────────

def test_edit_model_is_the_word_paragraphs(client, doc, people):
    m = _model(client, doc, people["editor"])
    assert m["mode"] == "docx" and m["editable"]
    assert [p["text"] for p in m["paragraphs"]] == [t for _, t, _ in PARAS]
    assert any(r["bold"] and r["text"] == "Effective Date" for r in m["paragraphs"][2]["runs"])


def test_outsider_on_open_matter_can_view_not_edit(client, doc, people):
    m = _model(client, doc, people["outsider"])
    assert not m["editable"]
    r = _save(client, doc, people["outsider"], m["base_version_id"], [{"op": "delete", "pid": 6}])
    assert r.status_code == 403
    assert client.post(f"/api/editor/documents/{doc}/lock", headers=as_member(people["outsider"])).status_code == 403


# ── locks ────────────────────────────────────────────────────────────────────

def test_one_editor_at_a_time(client, doc, people):
    a, b = as_member(people["editor"]), as_member(people["editor2"])
    assert client.post(f"/api/editor/documents/{doc}/lock", headers=a).status_code == 200
    held = client.post(f"/api/editor/documents/{doc}/lock", headers=b)
    assert held.status_code == 409 and held.json()["detail"]["lock"]["member_id"] == people["editor"]
    base = _model(client, doc, people["editor2"])["base_version_id"]
    assert _save(client, doc, people["editor2"], base, [{"op": "delete", "pid": 6}]).status_code == 409
    assert client.post(f"/api/editor/documents/{doc}/lock/heartbeat", headers=a).status_code == 200
    assert client.delete(f"/api/editor/documents/{doc}/lock", headers=a).status_code == 204
    assert client.post(f"/api/editor/documents/{doc}/lock", headers=b).status_code == 200


def test_expired_lock_does_not_block(client, doc, people):
    client.post(f"/api/editor/documents/{doc}/lock", headers=as_member(people["editor"]))
    with connect() as conn:
        conn.execute("UPDATE document_locks SET expires_at = now() - interval '1 minute' WHERE document_id = %s", (doc,))
        conn.commit()
    assert client.post(f"/api/editor/documents/{doc}/lock", headers=as_member(people["editor2"])).status_code == 200


def _lock(client, doc, member, token=None, takeover=False):
    headers = {**as_member(member), **({"X-Edit-Lock": token} if token else {})}
    return client.post(f"/api/editor/documents/{doc}/lock" + ("?takeover=true" if takeover else ""), headers=headers)


def test_second_window_of_the_same_person_takes_over_explicitly(client, doc, people):
    me = people["editor"]
    first = _lock(client, doc, me)
    assert first.status_code == 200
    t1 = first.json()["lock_token"]
    # Reloading the same window keeps the lock and the token.
    again = _lock(client, doc, me, token=t1)
    assert again.status_code == 200 and again.json()["lock_token"] == t1
    # A second window is told, and can take over.
    other = _lock(client, doc, me)
    assert other.status_code == 409 and other.json()["detail"]["reason"] == "held_by_you_elsewhere"
    t2 = _lock(client, doc, me, takeover=True).json()["lock_token"]
    assert t2 != t1
    # The first window is now superseded everywhere it writes; the draft belongs to the new window.
    old = {**as_member(me), "X-Edit-Lock": t1}
    hb = client.post(f"/api/editor/documents/{doc}/lock/heartbeat", headers=old)
    assert hb.status_code == 409 and hb.json()["detail"]["reason"] == "superseded"
    base = _model(client, doc, me)["base_version_id"]
    assert client.put(f"/api/editor/documents/{doc}/draft", headers=old,
                      json={"base_version_id": base, "ops": [{"op": "delete", "pid": 6}]}).status_code == 409
    assert client.post(f"/api/editor/documents/{doc}/save", headers=old,
                       json={"base_version_id": base, "ops": [{"op": "delete", "pid": 6}]}).status_code == 409
    # Closing the old window releases nothing.
    assert client.delete(f"/api/editor/documents/{doc}/lock", headers=old).status_code == 204
    assert _model(client, doc, me)["lock"]["member_id"] == me
    new = {**as_member(me), "X-Edit-Lock": t2}
    r = client.post(f"/api/editor/documents/{doc}/save", headers=new,
                    json={"base_version_id": base, "ops": [{"op": "delete", "pid": 6}]})
    assert r.status_code == 201, r.text


def test_only_a_matter_manager_takes_over_from_someone_else(client, doc, people):
    from app import access

    holder, other = people["editor"], people["editor2"]
    with connect() as conn:
        level = access.matter_level(conn, other, people["matter_id"])
    _lock(client, doc, holder)
    blocked = _lock(client, doc, other)
    assert blocked.status_code == 409 and blocked.json()["detail"]["can_take_over"] == (level == "manage")
    r = _lock(client, doc, other, takeover=True)
    if level == "manage":
        assert r.status_code == 200
        with connect() as conn:
            ev = conn.execute("SELECT detail FROM document_events WHERE document_id = %s AND action = 'lock.takeover' "
                              "ORDER BY seq DESC LIMIT 1", (doc,)).fetchone()
        assert ev["detail"]["from_member_id"] == holder
        # The holder's draft is theirs to keep: someone else's takeover does not block it.
        base = _model(client, doc, holder)["base_version_id"]
        assert client.put(f"/api/editor/documents/{doc}/draft", headers=as_member(holder),
                          json={"base_version_id": base, "ops": [{"op": "delete", "pid": 6}]}).status_code == 200
    else:
        assert r.status_code == 403
    # A reader who cannot edit can never take over.
    assert _lock(client, doc, people["outsider"], takeover=True).status_code == 403


# ── save ─────────────────────────────────────────────────────────────────────

OPS = [
    {"op": "replace", "pid": 3, "text": "Either Party may terminate on sixty (60) days' written notice."},
    {"op": "delete", "pid": 6},
    {"op": "insert_after", "pid": 5, "text": "Disputed amounts may be withheld in good faith."},
]


def test_tracked_save_writes_word_revisions_by_the_editor(client, doc, people):
    base = _model(client, doc, people["editor"])["base_version_id"]
    original = _docx()
    r = _save(client, doc, people["editor"], base, OPS, note="Longer notice; interest removed")
    assert r.status_code == 201, r.text
    assert r.json()["version_number"] == 2

    new = _file(doc)
    expected = [t for _, t, _ in PARAS]
    expected[3] = OPS[0]["text"]
    expected.insert(6, OPS[2]["text"])
    del expected[7]
    assert view(new, accept=True) == expected
    assert view(new, accept=False) == [t for _, t, _ in PARAS]
    with connect() as conn:
        name = conn.execute("SELECT name FROM members WHERE member_id = %s", (people["editor"],)).fetchone()["name"]
        v = conn.execute("SELECT author_name, created_by_member_id, origin, change_summary FROM document_versions "
                         "WHERE document_id = %s AND version_number = 2", (doc,)).fetchone()
    assert (v["author_name"], v["created_by_member_id"], v["origin"]) == (name, people["editor"], "editor")
    assert v["change_summary"] == "Longer notice; interest removed"
    body = Document(io.BytesIO(new)).element.body
    authors = {el.get(qn("w:author")) for tag in ("w:ins", "w:del") for el in body.iter(qn(tag))}
    assert authors == {name}

    # Paragraphs nobody edited keep their exact formatting.
    old_p = [p._p for p in Document(io.BytesIO(original)).paragraphs]
    new_p = [p._p for p in Document(io.BytesIO(new)).paragraphs]
    for i in (0, 1, 2, 4, 5):
        assert formatting_signature(old_p[i]) == formatting_signature(new_p[i])


def test_saved_text_is_what_search_indexes(client, doc, people):
    base = _model(client, doc, people["editor"])["base_version_id"]
    assert _save(client, doc, people["editor"], base, OPS).status_code == 201
    with connect() as conn:
        text = " ".join(r["text"] for r in conn.execute(
            "SELECT c.text FROM chunks c JOIN documents d ON d.current_version_id = c.version_id WHERE d.document_id = %s", (doc,)))
    assert "sixty (60)" in text and "withheld in good faith" in text
    assert "two per cent (2%)" not in text


def test_save_right_after_upload_keeps_chunks_and_gets_every_vector(client, doc, people):
    """The upload's vectors may still be being written in the background when the save
    replaces the chunks: the new version must keep its chunks and end up fully embedded."""
    from app.documents.editing import wait_for_indexing

    r = client.post(f"/api/editor/documents/{doc}/versions", headers=as_member(people["editor"]),
                    files={"file": ("again.docx", _docx(), "application/octet-stream")})
    assert r.status_code == 201, r.text
    r = _save(client, doc, people["editor"], r.json()["version_id"], OPS)
    assert r.status_code == 201, r.text
    vid = r.json()["version_id"]
    with connect() as conn:
        total = conn.execute("SELECT count(*) AS n FROM chunks WHERE version_id = %s", (vid,)).fetchone()["n"]
    assert total > 0
    wait_for_indexing()
    with connect() as conn:
        row = conn.execute("SELECT count(embedding) AS done, count(*) AS n FROM chunks WHERE version_id = %s",
                           (vid,)).fetchone()
    assert row["n"] == total and row["done"] == total


def test_next_edit_starts_from_the_accepted_text(client, doc, people):
    base = _model(client, doc, people["editor"])["base_version_id"]
    _save(client, doc, people["editor"], base, OPS)
    m = _model(client, doc, people["editor"])
    assert m["version_number"] == 2 and m["has_revisions"]
    assert m["paragraphs"][3]["text"] == OPS[0]["text"]


def test_clean_save_has_no_revisions(client, doc, people):
    base = _model(client, doc, people["editor"])["base_version_id"]
    assert _save(client, doc, people["editor"], base, OPS, mode="clean").status_code == 201
    new = _file(doc)
    assert view(new, accept=True) == view(new, accept=False)
    assert b"<w:ins " not in new


def test_stale_base_is_refused(client, doc, people):
    base = _model(client, doc, people["editor"])["base_version_id"]
    assert _save(client, doc, people["editor"], base, OPS[:1]).status_code == 201
    stale = _save(client, doc, people["editor"], base, [{"op": "delete", "pid": 2}])
    assert stale.status_code == 409 and "newer version" in stale.json()["detail"]["message"]


def test_bad_paragraph_id_is_refused(client, doc, people):
    base = _model(client, doc, people["editor"])["base_version_id"]
    assert _save(client, doc, people["editor"], base, [{"op": "replace", "pid": 999, "text": "x"}]).status_code == 422


# ── drafts, upload, compare, history ─────────────────────────────────────────

def test_draft_autosave_round_trip(client, doc, people):
    h = as_member(people["editor"])
    base = _model(client, doc, people["editor"])["base_version_id"]
    assert client.put(f"/api/editor/documents/{doc}/draft", headers=h, json={"base_version_id": base, "ops": OPS[:1]}).status_code == 200
    assert _model(client, doc, people["editor"])["draft"]["ops"][0]["pid"] == 3
    assert _model(client, doc, people["editor2"])["draft"] is None, "drafts are private to their author"
    assert client.delete(f"/api/editor/documents/{doc}/draft", headers=h).status_code == 204


def test_uploaded_version_is_attributed_to_the_session_member(client, doc, people):
    base = _model(client, doc, people["editor"])["base_version_id"]
    r = client.post(f"/api/editor/documents/{doc}/versions", headers=as_member(people["editor2"]),
                    files={"file": ("counterparty markup.docx", _docx(), "application/octet-stream")},
                    data={"base_version_id": base, "note": "Counterparty mark-up", "author_name": "Someone Else"})
    assert r.status_code == 201, r.text
    with connect() as conn:
        v = conn.execute("SELECT created_by_member_id, origin FROM document_versions WHERE version_id = %s",
                         (r.json()["version_id"],)).fetchone()
    assert v["created_by_member_id"] == people["editor2"] and v["origin"] == "upload"


def test_compare_on_screen_and_as_tracked_word_file(client, doc, people):
    h = as_member(people["editor"])
    base = _model(client, doc, people["editor"])["base_version_id"]
    new_id = _save(client, doc, people["editor"], base, OPS).json()["version_id"]
    cmp = client.get(f"/api/editor/documents/{doc}/compare", params={"from": base, "to": new_id}, headers=h).json()
    assert cmp["stats"] == {"inserted": 1, "deleted": 1, "changed": 1, "unchanged": 5}
    changed = next(b for b in cmp["blocks"] if b["op"] == "replace")
    assert {"t": "ins", "text": "sixty"} in changed["segments"] or any(s["t"] == "ins" and "sixty" in s["text"] for s in changed["segments"])
    red = client.get(f"/api/editor/documents/{doc}/compare.docx", params={"from": base, "to": new_id}, headers=h)
    assert red.status_code == 200 and b"PK" == red.content[:2]
    assert view(red.content, accept=True) == view(_file(doc), accept=True)


def test_history_lists_who_did_what(client, doc, people):
    h = as_member(people["editor"])
    client.post(f"/api/editor/documents/{doc}/lock", headers=h)
    base = _model(client, doc, people["editor"])["base_version_id"]
    _save(client, doc, people["editor"], base, OPS[:1], note="notice")
    items = client.get(f"/api/editor/documents/{doc}/history", headers=h).json()["items"]
    assert [i["action"] for i in items][:2] == ["edit.save", "lock"]
    assert items[0]["member_id"] == people["editor"]


# ── formatting (plan 17, G1) ─────────────────────────────────────────────────

def test_formatting_changes_save_as_tracked_word_formatting(client, doc, people):
    m = _model(client, doc, people["editor"])
    assert {"Normal", "Title", "Heading 1"} <= set(m["styles"])
    fees = m["paragraphs"][5]
    assert fees["text"].startswith("The Client shall pay")
    runs = [{"text": "The "}, {"text": "Client", "bold": True}, {"text": " shall pay the "}, {"text": "Fees", "bold": True},
            {"text": " within "}, {"text": "forty-five (45) days", "italic": True}, {"text": " of invoice."}]
    ops = [
        {"op": "format", "pid": 5, "runs": runs},
        {"op": "format", "pid": 6, "style": "Heading 2"},
        {"op": "insert_after", "pid": 6, "text": "Schedule", "style": "Heading 1", "runs": [{"text": "Schedule", "underline": True}]},
    ]
    r = _save(client, doc, people["editor"], m["base_version_id"], ops)
    assert r.status_code == 201, r.text
    assert r.json()["stats"]["restyled"] >= 1 and r.json()["stats"]["format_skipped"] == 0
    after = _model(client, doc, people["editor"])
    p5 = after["paragraphs"][5]
    assert p5["text"] == fees["text"]
    assert any(x["text"] == "forty-five (45) days" and x["italic"] for x in p5["runs"])
    assert after["paragraphs"][6]["style"] == "Heading 2"
    assert after["paragraphs"][7]["style"] == "Heading 1" and after["paragraphs"][7]["runs"][0]["underline"]
    # The stored file records who formatted what.
    import zipfile

    with zipfile.ZipFile(io.BytesIO(_file(doc))) as z:
        xml = z.read("word/document.xml").decode()
    assert "w:rPrChange" in xml and "w:pPrChange" in xml


def test_formatting_ops_are_validated(client, doc, people):
    base = _model(client, doc, people["editor"])["base_version_id"]
    bad_runs = [{"op": "replace", "pid": 3, "text": "abc", "runs": [{"text": "xyz"}]}]
    assert _save(client, doc, people["editor"], base, bad_runs).status_code == 422
    assert _save(client, doc, people["editor"], base, [{"op": "format", "pid": 3}]).status_code == 422
    assert _save(client, doc, people["editor"], base, [{"op": "format", "pid": 3, "style": "No Such Style"}]).status_code == 422
