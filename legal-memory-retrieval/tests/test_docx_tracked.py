"""Tracked changes written into the original .docx: formatting kept, accept/reject views exact."""
from __future__ import annotations

import io

from docx import Document
from docx.oxml.ns import qn

from app.drafting.docx_tracked import apply_tracked_changes, view


def _docx() -> bytes:
    d = Document()
    p = d.add_paragraph()
    p.add_run("7.3 ").bold = True
    p.add_run("The Supplier shall use ")
    p.add_run("Good Industry Practice").italic = True
    p.add_run(" at all times.")
    d.add_paragraph("7.4 Notices must be in writing.")
    tab = d.add_paragraph("8.1 Amount")
    tab.runs[0]._r.append(tab.runs[0]._r.makeelement(qn("w:tab"), {}))
    d.add_paragraph("8.2 Last clause.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_replace_keeps_formatting_and_both_views_are_exact():
    src = _docx()
    out, stats = apply_tracked_changes(src, [
        {"op": "replace", "pid": 0, "text": "7.3 The Vendor shall use Good Industry Practice at all times."}])
    assert stats["replaced"] == 1
    assert view(out, accept=True)[0] == "7.3 The Vendor shall use Good Industry Practice at all times."
    assert view(out, accept=False)[0] == "7.3 The Supplier shall use Good Industry Practice at all times."
    p = Document(io.BytesIO(out)).paragraphs[0]._p
    runs = [(r.find(qn("w:rPr")), "".join(t.text for t in r.iter(qn("w:t")))) for r in p.iter(qn("w:r"))]
    assert any(t == "7.3 " and rpr is not None and rpr.find(qn("w:b")) is not None for rpr, t in runs)
    assert any("Good Industry Practice" in t and rpr is not None and rpr.find(qn("w:i")) is not None for rpr, t in runs)
    assert p.find(f".//{qn('w:del')}") is not None and p.find(f".//{qn('w:ins')}") is not None


def test_edit_across_a_formatting_boundary():
    out, _ = apply_tracked_changes(_docx(), [{"op": "replace", "pid": 0, "text": "7.3 The Supplier shall use best practice at all times."}])
    assert view(out, accept=True)[0] == "7.3 The Supplier shall use best practice at all times."
    assert view(out, accept=False)[0] == "7.3 The Supplier shall use Good Industry Practice at all times."


def test_insert_and_delete_mark_paragraphs_and_keep_order():
    out, stats = apply_tracked_changes(_docx(), [
        {"op": "insert_after", "pid": 1, "text": "7.5 New clause."}, {"op": "delete", "pid": 3}])
    assert stats == {"replaced": 0, "replaced_whole": 0, "deleted": 1, "inserted": 1, "rejected": 0}
    accepted = view(out, accept=True)
    assert accepted[1:3] == ["7.4 Notices must be in writing.", "7.5 New clause."] and "8.2 Last clause." not in accepted
    assert view(out, accept=False)[-1] == "8.2 Last clause." and "7.5 New clause." not in view(out, accept=False)


def test_complex_paragraph_is_replaced_whole_not_diffed():
    out, stats = apply_tracked_changes(_docx(), [{"op": "replace", "pid": 2, "text": "8.1 Amount payable"}])
    assert stats["replaced_whole"] == 1
    assert view(out, accept=True)[2] == "8.1 Amount payable"


def test_bad_ids_are_rejected():
    _, stats = apply_tracked_changes(_docx(), [{"op": "replace", "pid": 99, "text": "x"}, {"op": "move", "pid": 0}])
    assert stats["rejected"] == 2


def test_accept_all_gives_the_clean_new_version_with_formatting():
    from app.drafting.docx_tracked import accept_all

    out, _ = apply_tracked_changes(_docx(), [
        {"op": "replace", "pid": 0, "text": "7.3 The Vendor shall use Good Industry Practice at all times."},
        {"op": "insert_after", "pid": 1, "text": "7.5 New clause."}, {"op": "delete", "pid": 3}])
    clean = accept_all(out)
    doc = Document(io.BytesIO(clean))
    assert [p.text for p in doc.paragraphs] == [
        "7.3 The Vendor shall use Good Industry Practice at all times.", "7.4 Notices must be in writing.",
        "7.5 New clause.", "8.1 Amount\t"]
    assert doc.element.body.find(f".//{qn('w:ins')}") is None and doc.element.body.find(f".//{qn('w:del')}") is None
    assert doc.paragraphs[0].runs[0].bold and doc.paragraphs[0].runs[0].text == "7.3 "
