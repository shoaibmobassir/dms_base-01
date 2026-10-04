"""Documents as commits: a clean current version, a commit log, diff, restore, blame, Word import (plan 21, C2).

Each test works on the throwaway Word document from test_document_editor (a real open matter, deleted afterwards).
"""
from __future__ import annotations

import io

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.db.connection import connect
from app.documents import docx_review, editing
from app.drafting.docx_tracked import apply_tracked_changes, view
from tests.conftest import as_member
from tests.test_document_editor import OPS, PARAS, _docx, _file, _model, _save, doc, people  # noqa: F401  (fixtures)

ORIGINAL = [t for _, t, _ in PARAS]


def _commits(client, doc_id, member):
    r = client.get(f"/api/editor/documents/{doc_id}/commits", headers=as_member(member))
    assert r.status_code == 200, r.text
    return r.json()["items"]


def _name(member_id: str) -> str:
    with connect() as conn:
        return conn.execute("SELECT name FROM members WHERE member_id = %s", (member_id,)).fetchone()["name"]


def _has_markup(data: bytes) -> bool:
    body = Document(io.BytesIO(data)).element.body
    return bool(list(body.iter(qn("w:ins"))) or list(body.iter(qn("w:del"))))


def _chunk_text(doc_id: str) -> str:
    editing.wait_for_indexing()
    with connect() as conn:
        return " ".join(r["text"] for r in conn.execute(
            "SELECT c.text FROM chunks c JOIN documents d ON d.current_version_id = c.version_id WHERE d.document_id = %s",
            (doc_id,)))


# ── commit log ───────────────────────────────────────────────────────────────

def test_commit_log_lists_each_version_with_message_author_kind_and_size_change(client, doc, people):
    base = _model(client, doc, people["editor"])["base_version_id"]
    r = _save(client, doc, people["editor"], base, OPS, note="Longer notice; interest removed")
    assert r.status_code == 201
    items = _commits(client, doc, people["editor"])
    assert [i["version_number"] for i in items] == [2, 1]
    head, first = items
    assert head["is_current"] and not first["is_current"]
    assert head["message"] == "Longer notice; interest removed" and head["kind"] == "editor"
    assert head["author_name"] == _name(people["editor"]) and head["is_clean"] is True
    assert head["parent_version_id"] == first["version_id"]
    assert isinstance(head["chars_delta"], int) and first["chars_delta"] is None


def test_outsider_can_read_the_log_but_not_restore(client, doc, people):
    items = _commits(client, doc, people["outsider"])
    assert items and items[0]["is_current"]
    r = client.post(f"/api/editor/documents/{doc}/restore", headers=as_member(people["outsider"]),
                    json={"version_id": items[0]["version_id"], "base_version_id": items[0]["version_id"]})
    assert r.status_code == 403


def test_diff_is_an_alias_of_compare_and_counts_words(client, doc, people):
    h = as_member(people["editor"])
    base = _model(client, doc, people["editor"])["base_version_id"]
    new_id = _save(client, doc, people["editor"], base, OPS).json()["version_id"]
    a = client.get(f"/api/editor/documents/{doc}/diff", params={"from": base, "to": new_id}, headers=h).json()
    b = client.get(f"/api/editor/documents/{doc}/compare", params={"from": base, "to": new_id}, headers=h).json()
    assert a == b and a["stats"]["words_added"] > 0 and a["stats"]["words_removed"] > 0


# ── restore ──────────────────────────────────────────────────────────────────

def test_restore_makes_an_old_version_current_as_a_new_version(client, doc, people):
    h = as_member(people["editor"])
    first = _model(client, doc, people["editor"])["base_version_id"]
    v2 = _save(client, doc, people["editor"], first, OPS, note="edit").json()
    r = client.post(f"/api/editor/documents/{doc}/restore", headers=h,
                    json={"version_id": first, "base_version_id": v2["version_id"], "note": "client changed their mind"})
    assert r.status_code == 201, r.text
    out = r.json()
    assert out["version_number"] == 3 and out["restored_version_number"] == 1

    # History is not rewritten: three versions, the restore is the newest and points at what it restored.
    items = _commits(client, doc, people["editor"])
    assert [i["version_number"] for i in items] == [3, 2, 1]
    assert items[0]["kind"] == "restore" and items[0]["restored_from_version_id"] == first
    assert items[0]["message"] == "Restored version 1: client changed their mind"
    assert items[0]["is_current"] and items[0]["is_clean"] is True
    with connect() as conn:
        body = {r["version_number"]: " ".join(r["body"].split()) for r in conn.execute(
            "SELECT version_number, body FROM document_versions WHERE document_id = %s", (doc,))}
    assert body[3] == body[1] and body[3] != body[2]

    # The current file reads as the original, with no markup; the edit model starts from it.
    assert view(_file(doc), accept=True) == ORIGINAL and not _has_markup(_file(doc))
    assert [p["text"] for p in _model(client, doc, people["editor"])["paragraphs"]] == ORIGINAL

    # Search follows: the restored text is what is indexed, not the text it replaced.
    text = _chunk_text(doc)
    assert "thirty (30) days" in text and "sixty (60)" not in text


def test_restore_refuses_a_stale_base_the_current_version_and_unknown_versions(client, doc, people):
    h = as_member(people["editor"])
    first = _model(client, doc, people["editor"])["base_version_id"]
    v2 = _save(client, doc, people["editor"], first, OPS).json()["version_id"]
    url = f"/api/editor/documents/{doc}/restore"
    stale = client.post(url, headers=h, json={"version_id": first, "base_version_id": first})
    assert stale.status_code == 409 and "newer version" in stale.json()["detail"]["message"]
    same = client.post(url, headers=h, json={"version_id": v2, "base_version_id": v2})
    assert same.status_code == 422 and "already the current" in same.json()["detail"]["message"]
    assert client.post(url, headers=h, json={"version_id": "VER-NOPE", "base_version_id": v2}).status_code == 404


def test_restore_respects_another_editors_lock(client, doc, people):
    a, b = as_member(people["editor"]), as_member(people["editor2"])
    first = _model(client, doc, people["editor"])["base_version_id"]
    v2 = _save(client, doc, people["editor"], first, OPS).json()["version_id"]
    assert client.post(f"/api/editor/documents/{doc}/lock", headers=a).status_code == 200
    r = client.post(f"/api/editor/documents/{doc}/restore", headers=b, json={"version_id": first, "base_version_id": v2})
    assert r.status_code == 409


# ── a Word file with tracked changes ─────────────────────────────────────────

def _reviewed_copy(reviewer: str, base_paragraphs: bool = True) -> bytes:
    """The document with one tracked change by ``reviewer``; ``base_paragraphs=False`` also changes its base text."""
    data = _docx()
    if not base_paragraphs:
        d = Document(io.BytesIO(data))
        d.paragraphs[0].insert_paragraph_before("PRIVILEGED AND CONFIDENTIAL")
        buf = io.BytesIO()
        d.save(buf)
        data = buf.getvalue()
    out, _ = apply_tracked_changes(
        data, [{"op": "replace", "pid": 3 + (0 if base_paragraphs else 1),
                "text": "Either Party may terminate on sixty (60) days' written notice."}], author=reviewer)
    return out


def _upload(client, doc_id, member, data, base=None):
    r = client.post(f"/api/editor/documents/{doc_id}/versions", headers=as_member(member),
                    data={"base_version_id": base} if base else {},
                    files={"file": ("reviewed.docx", data, "application/octet-stream")})
    assert r.status_code == 201, r.text
    return r.json()


def test_a_reviewed_copy_of_the_current_version_is_one_clean_commit_credited_to_the_reviewer(client, doc, people):
    raw = _reviewed_copy("Alice Reviewer")
    out = _upload(client, doc, people["editor"], raw, base=_model(client, doc, people["editor"])["base_version_id"])
    assert len(out["commits"]) == 1 and out["tracked_changes"] >= 1 and out["authors"] == ["Alice Reviewer"]

    head = _commits(client, doc, people["editor"])[0]
    assert head["author_name"] == "Alice Reviewer" and head["is_clean"] is True and head["has_source_file"] is True
    assert head["message"].startswith("Imported") and "Alice Reviewer" in head["message"]
    stored = _file(doc)
    assert not _has_markup(stored) and view(stored, accept=True)[3].endswith("sixty (60) days' written notice.")
    # The upload as it arrived is kept as evidence.
    with connect() as conn:
        raw_uri = conn.execute("SELECT source_storage_uri FROM document_versions WHERE version_id = %s",
                               (head["version_id"],)).fetchone()["source_storage_uri"]
    from app.storage.object_store import get_object_store

    assert docx_review.has_revisions(get_object_store().get(raw_uri))


def test_a_file_whose_base_text_differs_becomes_two_commits(client, doc, people):
    raw = _reviewed_copy("Bob Reviewer", base_paragraphs=False)
    out = _upload(client, doc, people["editor"], raw)
    assert len(out["commits"]) == 2
    items = _commits(client, doc, people["editor"])
    assert [i["version_number"] for i in items][:3] == [3, 2, 1]
    base_commit, final_commit = items[1], items[0]
    assert base_commit["kind"] == "import" and "before its" in base_commit["message"]
    assert final_commit["author_name"] == "Bob Reviewer" and final_commit["parent_version_id"] == base_commit["version_id"]
    assert not _has_markup(_file(doc))


def test_a_clean_upload_is_one_commit_with_no_source_copy(client, doc, people):
    out = _upload(client, doc, people["editor"], _docx())
    assert "commits" not in out
    head = _commits(client, doc, people["editor"])[0]
    assert head["is_clean"] is True and head["has_source_file"] is False and head["kind"] == "upload"


# ── an old version that still carries someone else's pending changes ─────────

def test_saving_over_a_file_with_others_pending_changes_does_not_accept_them(client, doc, people):
    """Legacy rows keep their markup until cleaned: accepting a reviewer's proposal as a side effect of an edit would
    decide for them."""
    from app.storage.object_store import get_object_store

    tracked = _reviewed_copy("Carol Reviewer")
    from app.documents.editing import _store_version_file

    with connect() as conn:
        d = editing._document(conn, doc)
        uri = _store_version_file(d, 1, "legacy.docx", tracked, editing.DOCX_MIME)
        conn.execute("UPDATE document_versions SET storage_uri = %s WHERE version_id = %s", (uri, d["current_version_id"]))
        conn.commit()
    base = _model(client, doc, people["editor"])["base_version_id"]
    r = _save(client, doc, people["editor"], base, [{"op": "replace", "pid": 0, "text": "Master Services Agreement"}])
    assert r.status_code == 201, r.text
    assert r.json()["mode"] != "clean"
    assert docx_review.has_revisions(_file(doc))
    assert _commits(client, doc, people["editor"])[0]["is_clean"] is False


# ── blame ────────────────────────────────────────────────────────────────────

def test_blame_credits_each_paragraph_to_the_version_that_last_changed_it(client, doc, people):
    h = as_member(people["editor"])
    first = _model(client, doc, people["editor"])["base_version_id"]
    v2 = _save(client, doc, people["editor"], first, OPS, note="notice and interest").json()
    out = client.get(f"/api/editor/documents/{doc}/blame", headers=h)
    assert out.status_code == 200, out.text
    paras = out.json()["paragraphs"]
    by_text = {p["text"]: p for p in paras}
    changed = by_text[OPS[0]["text"]]
    untouched = by_text[ORIGINAL[0]]
    assert changed["version_number"] == 2 and changed["version_id"] == v2["version_id"]
    assert changed["author"] == _name(people["editor"]) and changed["message"] == "notice and interest"
    assert untouched["version_number"] == 1 and untouched["before_window"] is False

    # An earlier version can be blamed too: nothing there is credited to the later edit.
    old = client.get(f"/api/editor/documents/{doc}/blame", params={"version_id": first}, headers=h).json()["paragraphs"]
    assert {p["version_number"] for p in old} == {1}


def test_blame_names_the_restore_for_paragraphs_it_brought_back(client, doc, people):
    h = as_member(people["editor"])
    first = _model(client, doc, people["editor"])["base_version_id"]
    v2 = _save(client, doc, people["editor"], first, OPS).json()["version_id"]
    assert client.post(f"/api/editor/documents/{doc}/restore", headers=h,
                       json={"version_id": first, "base_version_id": v2}).status_code == 201
    paras = client.get(f"/api/editor/documents/{doc}/blame", headers=h).json()["paragraphs"]
    brought_back = next(p for p in paras if p["text"] == ORIGINAL[3])
    assert brought_back["version_number"] == 3 and brought_back["message"] == "Restored version 1"


# ── backfill of versions saved before the commit model ───────────────────────

def test_backfill_gives_an_old_tracked_version_a_clean_file_and_keeps_the_original(client, doc, people):
    from app.documents.clean_versions import clean_versions
    from app.documents.editing import _store_version_file
    from app.storage.object_store import get_object_store

    tracked = _reviewed_copy("Dana Reviewer")
    with connect() as conn:
        d = editing._document(conn, doc)
        old_uri = _store_version_file(d, 1, "legacy.docx", tracked, editing.DOCX_MIME)
        conn.execute("UPDATE document_versions SET storage_uri = %s, is_clean = NULL WHERE version_id = %s",
                     (old_uri, d["current_version_id"]))
        conn.commit()

        dry = clean_versions(conn, apply=False, document_id=doc)
        assert dry["to_clean"] == 1 and dry["applied"] is False
        assert conn.execute("SELECT storage_uri FROM document_versions WHERE version_id = %s",
                            (d["current_version_id"],)).fetchone()["storage_uri"] == old_uri, "a dry run changes nothing"

        done = clean_versions(conn, apply=True, document_id=doc)
        assert done["to_clean"] == 1
        row = conn.execute("SELECT storage_uri, source_storage_uri, is_clean FROM document_versions WHERE version_id = %s",
                           (d["current_version_id"],)).fetchone()
    assert row["is_clean"] is True and row["source_storage_uri"] == old_uri and row["storage_uri"] != old_uri
    store = get_object_store()
    assert not _has_markup(store.get(row["storage_uri"])) and docx_review.has_revisions(store.get(old_uri))
    assert view(store.get(row["storage_uri"]), accept=True)[3].endswith("sixty (60) days' written notice.")

    with connect() as conn:  # safe to run again: nothing left to do
        again = clean_versions(conn, apply=True, document_id=doc)
    assert again["checked"] == 0 and again.get("to_clean", 0) == 0


# ── who worked on the document ───────────────────────────────────────────────

def test_contributors_still_count_an_editors_words_after_a_clean_save(client, doc, people):
    """The stored file is clean, but the record of who changed what is still made from the save."""
    base = _model(client, doc, people["editor"])["base_version_id"]
    assert _save(client, doc, people["editor"], base, OPS).status_code == 201
    out = client.get(f"/api/editor/documents/{doc}/contributors", headers=as_member(people["editor"])).json()
    me = next(p for p in out["people"] if p["member_id"] == people["editor"])
    assert me["word"]["insertions"] >= 1 and me["word"]["deletions"] >= 1 and me["precentis"].get("edit.save") == 1


def test_an_import_credits_the_reviewers_in_the_contributors_list(client, doc, people):
    _upload(client, doc, people["editor"], _reviewed_copy("Erin Reviewer"))
    out = client.get(f"/api/editor/documents/{doc}/contributors", headers=as_member(people["editor"])).json()
    erin = next(p for p in out["people"] if p["name"] == "Erin Reviewer")
    assert erin["external"] and erin["word"]["insertions"] >= 1
