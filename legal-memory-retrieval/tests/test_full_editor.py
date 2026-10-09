"""The full Word editor's save (plan 22, W2b): the browser sends the edited .docx; the server decides who made each new
change, keeps the stored version clean, and refuses stale or unauthorised saves."""
from __future__ import annotations

import io
import re
import zipfile

from app.db.connection import connect
from app.documents import docx_review
from app.documents.docx_restamp import restamp_new_revisions
from tests.conftest import as_member
from tests.test_document_editor import _file, _model, doc, people  # noqa: F401  (fixtures)

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _with_tracked_insert(data: bytes, author: str, words: str) -> bytes:
    """The file as a browser editor would save it: one tracked insertion at the end of the first paragraph."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            raw = zin.read(item.filename)
            if item.filename == "word/document.xml":
                xml = raw.decode()
                ins = (f'<w:ins w:id="9001" w:author="{author}" w:date="2020-01-01T00:00:00Z">'
                       f'<w:r><w:t xml:space="preserve"> {words}</w:t></w:r></w:ins>')
                xml = re.sub(r"(<w:p[ >].*?)(</w:p>)", lambda m: m.group(1) + ins + m.group(2), xml, count=1, flags=re.S)
                raw = xml.encode()
            zout.writestr(item, raw)
    return out.getvalue()


def _member_name(member_id):
    with connect() as conn:
        return conn.execute("SELECT name FROM members WHERE member_id = %s", (member_id,)).fetchone()["name"]


def test_new_changes_are_credited_to_the_signed_in_member_not_the_browser(client, doc, people):
    base = _model(client, doc, people["editor"])["base_version_id"]
    edited = _with_tracked_insert(_file(doc), "Spoofed Person", "subject to clause 9")
    r = client.post(f"/api/editor/documents/{doc}/save-docx", headers=as_member(people["editor"]),
                    files={"file": ("x.docx", edited, DOCX)}, data={"base_version_id": base, "note": "full editor"})
    assert r.status_code == 201, r.text
    assert r.json()["changes"] == 1 and r.json()["version_number"] == 2
    stored = _file(doc)
    assert not docx_review.has_revisions(stored)  # the stored version is the clean document
    from docx import Document

    assert "subject to clause 9" in "\n".join(p.text for p in Document(io.BytesIO(stored)).paragraphs)
    with connect() as conn:
        rows = conn.execute("SELECT author_name, member_id FROM document_revision_authors WHERE version_id = %s",
                            (r.json()["version_id"],)).fetchall()
    names = {x["author_name"] for x in rows}
    assert _member_name(people["editor"]) in names and "Spoofed Person" not in names
    commits = client.get(f"/api/editor/documents/{doc}/commits", headers=as_member(people["editor"])).json()["items"]
    assert commits[0]["message"] == "full editor" and commits[0]["kind"] == "editor"


def test_a_word_editor_save_is_timed_and_replaces_the_documents_chunks_in_one_batch(client, doc, people):
    from app.observability.metrics import WORKBENCH_SAVE_SECONDS

    def saves() -> float:
        return WORKBENCH_SAVE_SECONDS.labels(path="word_editor")._sum.get()

    before = saves()
    base = _model(client, doc, people["editor"])["base_version_id"]
    edited = _with_tracked_insert(_file(doc), "Anyone", "a batch-inserted clause")
    r = client.post(f"/api/editor/documents/{doc}/save-docx", headers=as_member(people["editor"]),
                    files={"file": ("x.docx", edited, DOCX)}, data={"base_version_id": base, "note": "timed"})
    assert r.status_code == 201, r.text
    assert saves() > before
    with connect() as conn:
        rows = conn.execute("SELECT version_id, chunk_index, text FROM chunks WHERE document_id = %s ORDER BY chunk_index",
                            (doc,)).fetchall()
        blocks = conn.execute("SELECT count(*) AS n, count(DISTINCT sequence) AS d FROM document_blocks WHERE version_id = %s",
                              (r.json()["version_id"],)).fetchone()
    assert rows and {x["version_id"] for x in rows} == {r.json()["version_id"]}  # only the current version is indexed
    assert [x["chunk_index"] for x in rows] == sorted({x["chunk_index"] for x in rows})  # no duplicates
    with connect() as conn:
        body = conn.execute("SELECT body FROM document_versions WHERE version_id = %s", (r.json()["version_id"],)).fetchone()["body"]
    assert "a batch-inserted clause" in body
    assert blocks["n"] == blocks["d"] > 0  # one block per sequence, ids written back


def test_existing_changes_keep_their_author_and_new_ones_are_restamped():
    from docx import Document

    d = Document()
    d.add_paragraph("The Supplier shall deliver.")
    buf = io.BytesIO()
    d.save(buf)
    base = _with_tracked_insert(buf.getvalue(), "Ravi Kalra", "on time")
    saved = _with_tracked_insert(base, "Browser Claim", "and in full")
    out, n = restamp_new_revisions(base, saved, "Helena Voss")
    xml = zipfile.ZipFile(io.BytesIO(out)).read("word/document.xml").decode()
    assert n == 1
    assert 'w:author="Ravi Kalra"' in xml and 'w:author="Helena Voss"' in xml and "Browser Claim" not in xml


def test_stale_unauthorised_and_non_word_saves_are_refused(client, doc, people):
    base = _model(client, doc, people["editor"])["base_version_id"]
    edited = _with_tracked_insert(_file(doc), "x", "words")
    r = client.post(f"/api/editor/documents/{doc}/save-docx", headers=as_member(people["outsider"]),
                    files={"file": ("x.docx", edited, DOCX)}, data={"base_version_id": base})
    assert r.status_code in (403, 404)
    r = client.post(f"/api/editor/documents/{doc}/save-docx", headers=as_member(people["editor"]),
                    files={"file": ("x.docx", b"not a zip", DOCX)}, data={"base_version_id": base})
    assert r.status_code == 422
    ok = client.post(f"/api/editor/documents/{doc}/save-docx", headers=as_member(people["editor"]),
                     files={"file": ("x.docx", edited, DOCX)}, data={"base_version_id": base})
    assert ok.status_code == 201
    stale = client.post(f"/api/editor/documents/{doc}/save-docx", headers=as_member(people["editor"]),
                        files={"file": ("x.docx", edited, DOCX)}, data={"base_version_id": base})
    assert stale.status_code == 409


def test_autosaved_word_draft_is_private_current_only_and_cleared_by_a_save(client, doc, people):
    base = _model(client, doc, people["editor"])["base_version_id"]
    edited = _with_tracked_insert(_file(doc), "x", "draft words")
    h = as_member(people["editor"])
    put = client.put(f"/api/editor/documents/{doc}/draft-docx", headers=h, files={"file": ("d.docx", edited, DOCX)},
                     data={"base_version_id": base})
    assert put.status_code == 200, put.text
    info = client.get(f"/api/editor/documents/{doc}/draft-docx/info", headers=h).json()["draft"]
    assert info["current"] and info["size_bytes"] == len(edited)
    assert client.get(f"/api/editor/documents/{doc}/draft-docx", headers=h).content == edited
    # someone else sees no draft and cannot read it
    other = as_member(people["editor2"])
    assert client.get(f"/api/editor/documents/{doc}/draft-docx/info", headers=other).json()["draft"] is None
    assert client.get(f"/api/editor/documents/{doc}/draft-docx", headers=other).status_code == 404
    assert client.put(f"/api/editor/documents/{doc}/draft-docx", headers=as_member(people["outsider"]),
                      files={"file": ("d.docx", edited, DOCX)}, data={"base_version_id": base}).status_code in (403, 404)
    assert client.put(f"/api/editor/documents/{doc}/draft-docx", headers=h, files={"file": ("d.docx", b"nope", DOCX)},
                      data={"base_version_id": base}).status_code == 422
    # saving a version removes the draft (and its file)
    assert client.post(f"/api/editor/documents/{doc}/save-docx", headers=h, files={"file": ("x.docx", edited, DOCX)},
                       data={"base_version_id": base}).status_code == 201
    assert client.get(f"/api/editor/documents/{doc}/draft-docx/info", headers=h).json()["draft"] is None
    with connect() as conn:
        assert conn.execute("SELECT count(*) AS n FROM document_docx_drafts WHERE document_id = %s", (doc,)).fetchone()["n"] == 0


def test_a_draft_on_an_older_version_is_not_offered(client, doc, people):
    base = _model(client, doc, people["editor"])["base_version_id"]
    h = as_member(people["editor"])
    edited = _with_tracked_insert(_file(doc), "x", "late words")
    assert client.put(f"/api/editor/documents/{doc}/draft-docx", headers=h, files={"file": ("d.docx", edited, DOCX)},
                      data={"base_version_id": base}).status_code == 200
    # a colleague saves a newer version
    assert client.post(f"/api/editor/documents/{doc}/save-docx", headers=as_member(people["editor2"]),
                       files={"file": ("x.docx", _with_tracked_insert(_file(doc), "y", "newer"), DOCX)},
                       data={"base_version_id": base}).status_code == 201
    info = client.get(f"/api/editor/documents/{doc}/draft-docx/info", headers=h).json()["draft"]
    assert info is not None and info["current"] is False
    assert client.get(f"/api/editor/documents/{doc}/draft-docx", headers=h).status_code == 404
