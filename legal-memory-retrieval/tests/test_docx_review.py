"""Reading and resolving Word tracked changes, and Word comments both ways (plan 18) — unit level."""
from __future__ import annotations

import io

from docx import Document

from app.documents import docx_comments as C
from app.documents import docx_review as R
from tests.word_fixtures import build

FINAL = ["Services Agreement", "The Supplier shall deliver within 30 days", "The fee is payable monthly.",
         "Confidential information stays confidential.", "Termination on notice.", "Governing law is English law.",
         "Notices go to the registered office.", "", "Disputes go to arbitration.", "Payment is due on invoice."]
ORIGINAL = ["Services Agreement", "The Supplier shall deliver ", "The fee is USD 10,000 payable monthly.",
            "Confidential information stays confidential.", "Termination on notice.", "Notices go to the registered office.",
            "Payment is due on invoice.", "Disputes go to arbitration.", ""]


def _bold(data: bytes, pid: int, word: str) -> bool:
    p = Document(io.BytesIO(data)).paragraphs[pid]
    return any(r.bold for r in p.runs if word in r.text)


def test_every_revision_is_read_with_its_author_and_date():
    revs = R.read_revisions(build())
    assert [(r.type, r.author, r.pid) for r in revs] == [
        ("insert", "Ravi Kalra", 1), ("delete", "Trilegal", 2), ("format", "Kunal Lalit Kaistha", 3),
        ("paragraph_insert", "Trilegal", 5), ("insert", "Trilegal", 5), ("move_from", "Ravi Kalra", 7),
        ("move_to", "Ravi Kalra", 9)]
    assert all(r.date for r in revs)
    assert revs[2].detail == "Bold added"
    assert revs[1].author == "Trilegal"  # "Trilegal " in the file: one person


def test_final_and_original_views():
    d = build()
    assert R.text_view(d, "final") == FINAL
    assert R.text_view(d, "original") == ORIGINAL
    assert R.read_revisions(R.accept_everything(d)) == [] and R.read_revisions(R.reject_everything(d)) == []


def test_accept_or_reject_one_persons_changes():
    d = build()
    out, n = R.resolve(d, R.keys_by(d, authors=["Trilegal"]), accept=False)
    assert n == 3
    left = R.read_revisions(out)
    assert {r.author for r in left} == {"Ravi Kalra", "Kunal Lalit Kaistha"}
    final = R.text_view(out, "final")
    assert "The fee is USD 10,000 payable monthly." in final and "Governing law is English law." not in final
    # Others' changes are untouched and still resolvable.
    out2, _ = R.resolve(out, R.keys_by(out, authors=["Ravi Kalra"]), accept=True)
    assert R.text_view(out2, "final")[1] == "The Supplier shall deliver within 30 days"
    assert [r.author for r in R.read_revisions(out2)] == ["Kunal Lalit Kaistha"]


def test_formatting_changes_accept_and_reject():
    d = build()
    key = [r.key for r in R.read_revisions(d) if r.type == "format"]
    assert _bold(R.resolve(d, key, accept=True)[0], 3, "Confidential")
    assert not _bold(R.resolve(d, key, accept=False)[0], 3, "Confidential")


def test_accepting_person_by_person_equals_accepting_everything():
    d = build()
    cur = d
    for a in ("Kunal Lalit Kaistha", "Ravi Kalra", "Trilegal"):
        cur, _ = R.resolve(cur, R.keys_by(cur, authors=[a]), accept=True)
    assert R.text_view(cur, "final") == R.text_view(R.accept_everything(d), "final") == FINAL


def test_grouped_changes_and_contributors():
    d = build()
    changes = R.group_changes(d)
    assert len(changes) == 6  # Trilegal's paragraph mark and its text are one row per type
    people = {p["author"]: p for p in R.contributors(d)}
    assert people["Ravi Kalra"]["insertions"] == 1 and people["Ravi Kalra"]["moves"] == 1 and people["Ravi Kalra"]["comments"] == 1
    assert people["Trilegal"]["deletions"] == 1 and people["Trilegal"]["replies"] == 1
    assert people["Kunal Lalit Kaistha"]["formats"] == 1
    assert R.paragraph_pending(d)[5][0]["author"] == "Trilegal"


def test_word_comments_are_read_as_threads():
    cs = C.read_comments(build())
    assert [(c["author"], c["quote"], bool(c["parent_para_id"]), c["done"]) for c in cs] == [
        ("Ravi Kalra", "within 30 days", False, False), ("Trilegal", "", True, False),
        ("Kunal Lalit Kaistha", "Confidential", False, True)]


def test_precentis_comments_are_written_into_the_file_and_back():
    d = build()
    root = C.read_comments(d)[0]
    threads = [
        {"annotation_id": "CMT-A", "parent_id": None, "source": "precentis", "external_id": None, "author": "Amina El-Sayed",
         "text": "Cap the fee?", "status": "open", "quote": "payable monthly", "pid": 2},
        {"annotation_id": "CMT-B", "parent_id": "CMT-A", "source": "precentis", "external_id": None, "author": "Helena Voss",
         "text": "Yes, at USD 12,000.", "status": "open"},
        {"annotation_id": "CMT-W", "parent_id": None, "source": "word", "external_id": C.external_id(root), "author": "Ravi Kalra",
         "text": root["text"], "status": "resolved"},
    ]
    out, written = C.write_comments(d, threads)
    back = {c["para_id"]: c for c in C.read_comments(out)}
    a, b = back[written["CMT-A"]], back[written["CMT-B"]]
    assert a["author"] == "Amina El-Sayed" and a["quote"] == "payable monthly" and a["pid"] == 2
    assert b["parent_para_id"] == written["CMT-A"] and b["text"] == "Yes, at USD 12,000."
    assert back[root["para_id"]]["done"] is True  # resolved in Precentis → done in Word
    assert R.read_revisions(out) == R.read_revisions(d)  # tracked changes untouched
    # Writing the same threads again adds nothing.
    again, more = C.write_comments(out, [{**t, "external_id": written.get(t["annotation_id"], t["external_id"])} for t in threads])
    assert more == {} and len(C.read_comments(again)) == len(C.read_comments(out))
    Document(io.BytesIO(out))  # still a valid package


def test_comments_are_added_to_a_file_without_any():
    doc = Document()
    doc.add_paragraph("A plain clause about delivery.")
    buf = io.BytesIO()
    doc.save(buf)
    out, written = C.write_comments(buf.getvalue(), [{"annotation_id": "CMT-X", "parent_id": None, "source": "precentis",
                                                      "external_id": None, "author": "Amina El-Sayed", "text": "Why plain?",
                                                      "status": "open", "quote": "about delivery", "pid": 0}])
    cs = C.read_comments(out)
    assert len(cs) == 1 and cs[0]["quote"] == "about delivery" and cs[0]["para_id"] == written["CMT-X"]
    Document(io.BytesIO(out))
