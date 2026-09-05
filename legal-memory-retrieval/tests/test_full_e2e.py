"""
Full End-to-End (E2E) Verification Suite for FirmOS / Legal DMS Intelligence Suite.
Tests the entire lawyer journey from institutional retrieval to matrix review, drafting,
Word add-in integration, case law citation verification, and tamper-evident packaging.
"""

import io
import json
import zipfile
import pytest
from fastapi.testclient import TestClient
import openpyxl

from app.api.main import app
from app.audit.manifest_signer import get_manifest_signer
from app.auth.handoff import get_auth_handoff_service
from app.auth.key_vault import KeyVault
from app.caselaw.citation_parser import get_citation_parser
from app.caselaw.courtlistener_client import get_courtlistener_client
from app.drafting.clause_diff_service import get_clause_diff_service
from app.drafting.docx_redline_generator import DocxRedlineGenerator
from app.review.spreadsheet_exporter import SpreadsheetExporter
from app.workflows.catalog_loader import get_catalog_loader
from app.workflows.engine import get_workflow_engine


@pytest.fixture
def client():
    return TestClient(app)


class TestFullEndToEndLawyerJourney:
    """E2E workflow testing across the unified Legal Operating System."""

    def test_01_system_discovery_and_health(self, client):
        """Verify service catalog and health discovery endpoints."""
        resp = client.get("/")
        assert resp.status_code == 200
        data = resp.json()
        assert data["service"] == "LEXOS Legal DMS"
        assert "services" in data
        assert "tabular" in data["services"]
        assert "workflows" in data["services"]
        assert "drafting" in data["services"]
        assert "word" in data["services"]
        assert "caselaw" in data["services"]
        assert "audit" in data["services"]

        # Check all health endpoints
        for svc_name, svc_info in data["services"].items():
            h_resp = client.get(svc_info["health"])
            assert h_resp.status_code == 200, f"Health check failed for {svc_name}"
            assert h_resp.json().get("status") == "ok"

    def test_02_key_vault_and_byok_security(self):
        """Verify tenant API keys encryption with AES-256-GCM and PBKDF2."""
        vault = KeyVault(master_secret="firmos-e2e-master-encryption-key-2026")
        tenant_key = "sk-ant-api03-firm-custom-claude-key-998877"
        encrypted_token = vault.encrypt_key(tenant_key)
        assert encrypted_token != tenant_key
        assert len(encrypted_token) > 32

        # Decrypt
        decrypted_key = vault.decrypt_key(encrypted_token)
        assert decrypted_key == tenant_key

        # Tampering check
        corrupted_payload = encrypted_token[:10] + "ZZZZ" + encrypted_token[14:]
        with pytest.raises(ValueError):
            vault.decrypt_key(corrupted_payload)

    def test_03_hybrid_retrieval_and_acl(self, client):
        """Verify 5-channel hybrid retrieval with permission filtering."""
        retrieve_payload = {
            "query": "Limitation of liability and indemnities in commercial contracts",
            "k": 5,
        }
        resp = client.post("/api/retrieval", json=retrieve_payload, headers={"X-Member-Id": "MEM-00001"})
        assert resp.status_code == 200
        data = resp.json()
        assert "hits" in data
        assert isinstance(data["hits"], list)
        assert "latency_ms" in data

    def test_04_ask_firm_ai_with_grounding(self, client):
        """Verify Ask Firm AI engine returns answer with citations or abstains gracefully."""
        ask_payload = {
            "query": "What are the firm's standard positions on warranty survival periods?",
            "k": 5,
        }
        resp = client.post("/api/answers", json=ask_payload, headers={"X-Member-Id": "MEM-00001"})
        assert resp.status_code == 200
        data = resp.json()
        assert "answer" in data
        assert "citations" in data
        assert "abstained" in data
        assert isinstance(data["citations"], list)

    def test_05_tabular_due_diligence_matrix_review(self, client):
        """E2E test of high-throughput Tabular Document Review, cell override, and Excel export."""
        create_payload = {
            "title": "M&A Due Diligence - NDA & Vendor Portfolio",
            "document_ids": ["DOC-00001", "DOC-00002", "DOC-00003"],
            "columns": [
                {
                    "id": "col_law",
                    "label": "Governing Law",
                    "prompt": "Identify the governing law and jurisdiction clause",
                    "data_type": "text",
                },
                {
                    "id": "col_cap",
                    "label": "Liability Cap",
                    "prompt": "Extract the financial limitation of liability cap amount",
                    "data_type": "currency",
                },
                {
                    "id": "col_indemnity",
                    "label": "Indemnity Scope",
                    "prompt": "Determine if indemnity covers third-party IP claims",
                    "data_type": "boolean",
                },
            ],
        }
        resp = client.post("/api/tabular/reviews", json=create_payload)
        assert resp.status_code == 200
        review_data = resp.json()
        review_id = review_data["review_id"]
        assert review_data["title"] == "M&A Due Diligence - NDA & Vendor Portfolio"
        assert len(review_data["document_ids"]) == 3
        assert "DOC-00001" in review_data["cells"]

        # Human lawyer overrides one cell
        patch_payload = {
            "value": "Delaware Law with New York Venue",
            "reasoning": "Counsel verified per Clause 18.2 amendment",
        }
        patch_resp = client.patch(
            f"/api/tabular/reviews/{review_id}/cells/DOC-00001/col_law",
            json=patch_payload,
        )
        assert patch_resp.status_code == 200
        patched_cell = patch_resp.json()
        assert patched_cell["value"] == "Delaware Law with New York Venue"
        assert patched_cell["is_overridden"] is True

        # Export to Excel (.xlsx)
        export_resp = client.get(f"/api/tabular/reviews/{review_id}/export/xlsx")
        assert export_resp.status_code == 200
        assert export_resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        
        # Verify valid Excel structure
        wb = openpyxl.load_workbook(io.BytesIO(export_resp.content))
        assert "Review Matrix" in wb.sheetnames
        ws = wb["Review Matrix"]
        assert ws["C4"].value is not None

    def test_06_declarative_legal_playbooks(self, client):
        """E2E execution of multi-step declarative YAML legal workflows."""
        # 1. List catalog
        list_resp = client.get("/api/workflows/")
        assert list_resp.status_code == 200
        wfs = list_resp.json()
        assert len(wfs) >= 3

        # 2. Run Cease & Desist Drafter
        cd_payload = {
            "workflow_id": "cease-and-desist-drafter",
            "inputs": {
                "rights_holder": "Lexis Innovations Inc.",
                "infringing_party": "Copycat Legal Systems Ltd.",
                "infringing_details": "Unauthorized scraping and duplication of proprietary knowledge base",
                "remedy_deadline_days": 10,
            },
        }
        cd_resp = client.post("/api/workflows/run", json=cd_payload)
        assert cd_resp.status_code == 200
        cd_run = cd_resp.json()
        assert cd_run["status"] == "completed"
        assert "LEGAL DEMAND" in cd_run["final_output"] or "Lexis Innovations" in cd_run["final_output"]

        # 3. Run Contract Triage Router
        triage_payload = {
            "workflow_id": "contract-triage-router",
            "inputs": {
                "contract_title": "Master Services Agreement v2.1",
                "governing_law": "English Law",
            },
        }
        triage_resp = client.post("/api/workflows/run", json=triage_payload)
        assert triage_resp.status_code == 200
        triage_run = triage_resp.json()
        assert triage_run["status"] == "completed"
        assert "CONTRACT TRIAGE" in triage_run["final_output"] or "Executive Summary" in triage_run["final_output"]

    def test_07_native_docx_track_changes_and_clause_diff(self, client):
        """E2E test of clause risk analysis and native Word OpenXML tracked changes (<w:ins>, <w:del>)."""
        # 1. Clause Risk Analysis
        clause_payload = {
            "clause_text": "Supplier shall defend and hold harmless Customer from any claims arising out of the Agreement without financial cap.",
            "clause_type": "Indemnification",
            "client_stance": "Supplier",
        }
        clause_resp = client.post("/api/drafting/redline-clause", json=clause_payload)
        assert clause_resp.status_code == 200
        clause_data = clause_resp.json()
        assert clause_data["risk_severity"] in ("low", "medium", "high", "critical")
        assert "proposed_redline" in clause_data
        assert "deviation_analysis" in clause_data

        # 2. Generate Tracked Changes DOCX
        docx_payload = {
            "original_text": "The aggregate liability of either party shall not exceed $50,000.",
            "revised_text": "The aggregate liability of either party shall not exceed the total fees paid in the preceding 12 months.",
            "title": "Liability Clause Redline",
            "author": "FirmOS Partner AI",
        }
        docx_resp = client.post("/api/drafting/generate-docx-tracked", json=docx_payload)
        assert docx_resp.status_code == 200
        assert docx_resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

        # Verify OpenXML elements in docx zip
        zf = zipfile.ZipFile(io.BytesIO(docx_resp.content))
        assert "word/document.xml" in zf.namelist()
        doc_xml = zf.read("word/document.xml").decode("utf-8")
        assert "w:ins" in doc_xml
        assert "w:del" in doc_xml

    def test_08_word_add_in_auth_and_selection_ingestion(self, client):
        """E2E test of browser-to-Word taskpane one-time auth ticket and selection analysis."""
        # 1. Generate one-time handoff ticket
        ticket_resp = client.post(
            "/api/word/handoff/ticket",
            json={"member_id": "MEM-00001", "email": "partner@apexchambers.com"},
        )
        assert ticket_resp.status_code == 200
        ticket = ticket_resp.json()["ticket"]
        assert ticket.startswith("TKT-")

        # 2. Exchange ticket inside Word add-in
        exchange_resp = client.post(
            "/api/word/handoff/exchange",
            json={"ticket": ticket},
        )
        assert exchange_resp.status_code == 200
        session_data = exchange_resp.json()
        assert session_data["member_id"] == "MEM-00001"
        assert session_data["email"] == "partner@apexchambers.com"

        # 3. Analyze selection in Word
        sel_resp = client.post(
            "/api/word/analyze-selection",
            json={
                "selection_text": "Each party shall treat confidential information with at least reasonable care.",
                "action": "explain",
            },
        )
        assert sel_resp.status_code == 200
        assert sel_resp.json()["action"] == "explain"
        assert "analysis" in sel_resp.json()

        # 4. Draft new clause for Word
        draft_resp = client.post(
            "/api/word/draft-clause",
            json={
                "instruction": "Draft a reciprocal non-disclosure clause with standard trade secret protections.",
                "clause_type": "Confidentiality",
            },
        )
        assert draft_resp.status_code == 200
        assert "clause_text" in draft_resp.json()

    def test_09_case_law_citation_parsing_and_verification(self, client):
        """E2E test of legal citation extraction and CourtListener precedential authority lookup."""
        brief_text = (
            "As established in Chevron U.S.A. Inc. v. Natural Resources Defense Council, Inc., "
            "467 U.S. 837 (1984), and confirmed in 987 F.3d 654, agencies receive deference. "
            "Furthermore, 28 U.S.C. § 1331 provides federal question jurisdiction."
        )
        extract_resp = client.post(
            "/api/caselaw/extract-citations",
            json={"text": brief_text},
        )
        assert extract_resp.status_code == 200
        citations = extract_resp.json()
        assert len(citations) >= 3
        normalized_cits = [c["normalized"] for c in citations]
        assert any("467 U.S. 837" in n for n in normalized_cits)

        # Verify authority
        verify_resp = client.post(
            "/api/caselaw/verify",
            json={"citation": "467 U.S. 837"},
        )
        assert verify_resp.status_code == 200
        opinion = verify_resp.json()
        assert opinion["citation"] == "467 U.S. 837"
        assert opinion["is_good_law"] is not None
        assert opinion["precedential_status"] in ("Precedential", "Published")

    def test_10_tamper_evident_signed_export_bundle(self, client):
        """E2E test of cryptographic manifest signing and tamper-evident verification."""
        export_payload = {
            "title": "Closing Trial Bundle & Review Findings",
            "matter_id": "MAT-00042",
            "files": [
                {
                    "path": "memorandum.md",
                    "content": "# Legal Opinion\nAll material liabilities have been disclosed.",
                },
                {
                    "path": "risk_matrix.csv",
                    "content": "contract_id,risk_level\nDOC-001,Low\nDOC-002,Medium\n",
                },
            ],
            "metadata": {"export_version": "v1.0", "classified": True},
        }
        export_resp = client.post("/api/audit/export-bundle", json=export_payload)
        assert export_resp.status_code == 200
        assert export_resp.headers["content-type"] == "application/zip"

        # Inspect ZIP bundle and MANIFEST.json
        zf = zipfile.ZipFile(io.BytesIO(export_resp.content))
        assert "MANIFEST.json" in zf.namelist()
        assert "memorandum.md" in zf.namelist()
        assert "risk_matrix.csv" in zf.namelist()

        manifest_str = zf.read("MANIFEST.json").decode("utf-8")
        manifest_data = json.loads(manifest_str)
        assert "hmac_sha256_signature" in manifest_data
        assert len(manifest_data["files"]) == 2

        # Verify authentic signature
        verify_resp = client.post(
            "/api/audit/verify-manifest",
            json={"manifest_json": manifest_str},
        )
        assert verify_resp.status_code == 200
        assert verify_resp.json()["is_valid"] is True
        assert verify_resp.json()["status"] == "AUTHENTIC"

        # Test tamper detection
        tampered_manifest = manifest_str.replace("Closing Trial Bundle & Review Findings", "Forged Trial Bundle & Review Findings")
        tampered_resp = client.post(
            "/api/audit/verify-manifest",
            json={"manifest_json": tampered_manifest},
        )
        assert tampered_resp.status_code == 200
        assert tampered_resp.json()["is_valid"] is False
        assert tampered_resp.json()["status"] == "TAMPERED_OR_INVALID"

