"""Word tracked *formatting* changes for the browser editor (plan 17, G1).

``app/drafting/docx_tracked.py`` writes the text changes (insertions and deletions); this
module runs on its output and writes what the editor asks about formatting:

  style   the paragraph style (Normal, Title, Heading 1, …) → ``w:pStyle`` plus
          ``w:pPrChange`` holding the previous paragraph properties
  runs    bold / italic / underline over the paragraph's text as it reads with the
          changes accepted → runs split at the boundaries, explicit ``w:b``/``w:i``/``w:u``,
          and ``w:rPrChange`` holding each changed run's previous properties

Word shows these as "Formatted: Bold" / "Formatted: Heading 1" under the author, and
rejecting them restores the old formatting. Text the same save inserts is already an
insertion, so it takes its formatting directly (no formatting revision on new text).

Paragraph ids are indexes into the *original* body paragraphs (as for the text ops):
after the text pass, the k-th paragraph that is not itself an insertion is pid k, and
paragraphs inserted after pid k follow it in the order of their ``insert_after`` ops.
"""
from __future__ import annotations

import copy
import io
from collections import defaultdict
from datetime import datetime, timezone

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.run import Run

_FLAGS = ("bold", "italic", "underline")
_SIMPLE = {qn("w:rPr"), qn("w:t")}


class _Marks:
    """Revision ids/author/date for the formatting revisions of one save."""

    def __init__(self, author: str, date: str | None, start: int = 20000):
        self.author = author
        self.date = date or datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        self.n = start

    def element(self, tag: str):
        self.n += 1
        el = OxmlElement(tag)
        el.set(qn("w:id"), str(self.n))
        el.set(qn("w:author"), self.author)
        el.set(qn("w:date"), self.date)
        return el


def has_format_revisions(docx_xml: bytes) -> bool:
    return b"w:rPrChange" in docx_xml or b"w:pPrChange" in docx_xml


def wants_formatting(op: dict) -> bool:
    return op.get("op") in ("replace", "insert_after", "format") and (op.get("style") is not None or op.get("runs") is not None)


def paragraph_styles(doc) -> dict[str, str]:
    """Paragraph style name → style id, for the styles this document defines."""
    return {s.name: s.style_id for s in doc.styles if s.type == WD_STYLE_TYPE.PARAGRAPH and s.name}


# ── locating paragraphs after the text pass ──────────────────────────────────

def _is_inserted(p_el) -> bool:
    ppr = p_el.find(qn("w:pPr"))
    rpr = ppr.find(qn("w:rPr")) if ppr is not None else None
    return rpr is not None and rpr.find(qn("w:ins")) is not None


TAG = "{urn:precentis:edit}pid"


def _targets(doc, ops: list[dict]):
    """(paragraph element, op, is_new_paragraph) for every op that asks for formatting.

    When the save tagged its paragraphs (``editing._tag_paragraphs``), the first paragraph with a
    tag is the original and later copies are this save's insertions after it — reliable even when
    the file holds other reviewers' inserted paragraphs. Otherwise paragraphs whose mark is an
    insertion are taken to be this save's."""
    originals, after = [], defaultdict(list)
    if any(p._p.get(TAG) is not None for p in doc.paragraphs):
        first: dict[int, object] = {}
        for p in doc.paragraphs:
            tag = p._p.get(TAG)
            if tag is None:
                continue
            pid = int(tag)
            if pid in first:
                after[pid].append(p._p)
            else:
                first[pid] = p._p
        originals = [first.get(i) for i in range(max(first) + 1)] if first else []
    else:
        for p in doc.paragraphs:
            if _is_inserted(p._p):
                if originals:
                    after[len(originals) - 1].append(p._p)
            else:
                originals.append(p._p)
    seen = defaultdict(int)  # insert_after ops per anchor, counted whether or not they format
    for op in ops:
        pid = op.get("pid")
        if op.get("op") == "insert_after":
            k = seen[pid]
            seen[pid] += 1
            if wants_formatting(op) and k < len(after.get(pid, [])):
                yield after[pid][k], op, True
        elif wants_formatting(op) and isinstance(pid, int) and 0 <= pid < len(originals) and originals[pid] is not None:
            yield originals[pid], op, False


# ── paragraph style ──────────────────────────────────────────────────────────

def _set_style(p_el, style_id: str, marks: _Marks | None) -> bool:
    ppr = p_el.get_or_add_pPr()
    current = ppr.find(qn("w:pStyle"))
    if (current.get(qn("w:val")) if current is not None else None) == style_id:
        return False
    before = [copy.deepcopy(c) for c in ppr if c.tag not in (qn("w:rPr"), qn("w:sectPr"), qn("w:pPrChange"))]
    if current is not None:
        ppr.remove(current)
    el = OxmlElement("w:pStyle")
    el.set(qn("w:val"), style_id)
    ppr.insert(0, el)
    if marks is not None and ppr.find(qn("w:pPrChange")) is None:  # keep the oldest recorded state
        change = marks.element("w:pPrChange")
        old = OxmlElement("w:pPr")
        for c in before:
            old.append(c)
        change.append(old)
        ppr.append(change)  # last child of pPr
    return True


# ── run formatting ───────────────────────────────────────────────────────────

def _accepted_runs(p_el):
    """(run element, inside an insertion) in reading order, skipping deleted text."""
    for child in p_el:
        if child.tag == qn("w:r"):
            yield child, False
        elif child.tag == qn("w:ins"):
            for r in child:
                if r.tag == qn("w:r"):
                    yield r, True
        elif child.tag == qn("w:hyperlink"):
            for r in child.iter(qn("w:r")):
                yield r, False


def _flags(run: Run) -> tuple[bool, bool, bool]:
    return bool(run.bold), bool(run.italic), bool(run.underline)


def _desired(runs: list[dict]) -> tuple[str, list[tuple[bool, bool, bool]]]:
    text, flags = "", []
    for r in runs:
        t = str(r.get("text") or "")
        text += t
        flags += [tuple(bool(r.get(f)) for f in _FLAGS)] * len(t)
    return text, flags


def _segments(flags: list[tuple], start: int, end: int):
    """Split [start, end) where the wanted formatting changes: (start, end, flags)."""
    s = start
    for i in range(start + 1, end + 1):
        if i == end or flags[i] != flags[s]:
            yield s, i, flags[s]
            s = i


def _split_run(r, pieces: list[tuple[int, int]]) -> list:
    """Replace a text-only run by copies holding the given slices of its text."""
    text = "".join(t.text or "" for t in r.iter(qn("w:t")))
    out = []
    for a, b in pieces:
        new = copy.deepcopy(r)
        ts = list(new.iter(qn("w:t")))
        for extra in ts[1:]:
            extra.getparent().remove(extra)
        ts[0].text = text[a:b]
        ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r.addprevious(new)
        out.append(new)
    r.getparent().remove(r)
    return out


def _apply_flags(r, want: tuple[bool, bool, bool], marks: _Marks | None) -> bool:
    run = Run(r, None)
    if _flags(run) == want:
        return False
    rpr = r.get_or_add_rPr()
    if marks is not None and rpr.find(qn("w:rPrChange")) is None:
        before = [copy.deepcopy(c) for c in rpr if c.tag != qn("w:rPrChange")]
        change = marks.element("w:rPrChange")
        old = OxmlElement("w:rPr")
        for c in before:
            old.append(c)
        change.append(old)
    else:
        change = None
    run.bold = True if want[0] else None
    run.italic = True if want[1] else None
    run.underline = True if want[2] else None
    rpr = r.get_or_add_rPr()
    existing = rpr.find(qn("w:rPrChange"))
    if existing is not None:  # keep it last (schema order)
        rpr.remove(existing)
        rpr.append(existing)
    elif change is not None:
        rpr.append(change)
    return True


def _format_runs(p_el, runs: list[dict], marks: _Marks, new_paragraph: bool) -> tuple[int, int]:
    """Returns (runs changed, runs that could not be changed)."""
    want_text, want = _desired(runs)
    items = list(_accepted_runs(p_el))
    texts = [Run(r, None).text for r, _ in items]
    if "".join(texts) != want_text:
        return 0, 1
    changed = skipped = 0
    pos = 0
    for (r, inserted), text in zip(items, texts):
        start, end = pos, pos + len(text)
        pos = end
        if start == end:
            continue
        segs = list(_segments(want, start, end))
        track = None if (inserted or new_paragraph) else marks
        if len(segs) == 1:
            changed += _apply_flags(r, segs[0][2], track)
            continue
        if not {c.tag for c in r} <= _SIMPLE:  # tabs, breaks, fields: do not split
            skipped += 1
            continue
        pieces = _split_run(r, [(a - start, b - start) for a, b, _ in segs])
        for piece, (_, _, flags) in zip(pieces, segs):
            changed += _apply_flags(piece, flags, track)
    return changed, skipped


# ── entry points ─────────────────────────────────────────────────────────────

def apply_formatting(docx_bytes: bytes, ops: list[dict], *, author: str, date: str | None = None) -> tuple[bytes, dict]:
    """Write the formatting the ops ask for into ``docx_bytes`` (the text pass's output)."""
    doc = Document(io.BytesIO(docx_bytes))
    styles = paragraph_styles(doc)
    marks = _Marks(author, date)
    stats = {"restyled": 0, "formatted_runs": 0, "format_skipped": 0}
    for p_el, op, new in list(_targets(doc, ops)):
        style = op.get("style")
        if style is not None:
            if style in styles:
                stats["restyled"] += _set_style(p_el, styles[style], None if new else marks)
            else:
                stats["format_skipped"] += 1
        if op.get("runs") is not None:
            changed, skipped = _format_runs(p_el, op["runs"], marks, new)
            stats["formatted_runs"] += changed
            stats["format_skipped"] += skipped
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), stats


def _body(data: bytes):
    doc = Document(io.BytesIO(data))
    return doc, doc.element.body


def accept_formatting(data: bytes) -> bytes:
    """The same file with every formatting revision accepted (the new formatting stays)."""
    doc, body = _body(data)
    for tag in ("w:rPrChange", "w:pPrChange"):
        for el in list(body.iter(qn(tag))):
            el.getparent().remove(el)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def reject_formatting(data: bytes) -> bytes:
    """The same file with every formatting revision rejected (the old formatting back)."""
    doc, body = _body(data)
    for change in list(body.iter(qn("w:rPrChange"))):
        rpr = change.getparent()
        old = change.find(qn("w:rPr"))
        for c in list(rpr):
            rpr.remove(c)
        for c in list(old if old is not None else []):
            rpr.append(c)
    for change in list(body.iter(qn("w:pPrChange"))):
        ppr = change.getparent()
        old = change.find(qn("w:pPr"))
        keep = [c for c in ppr if c.tag in (qn("w:rPr"), qn("w:sectPr"))]
        for c in list(ppr):
            ppr.remove(c)
        for c in list(old if old is not None else []):
            ppr.append(c)
        for c in keep:
            ppr.append(c)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
