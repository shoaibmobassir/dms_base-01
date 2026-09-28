"""The editable form of a firm document: its paragraphs, and the original Word file when there is one.

Paragraph ids used by the edit engine and by ``docx_tracked`` must name the same paragraphs, so
a .docx is read from its ORIGINAL file (``Document.paragraphs``: body paragraphs, tables
excluded). Other formats are edited as text: the extracted pages split into paragraphs.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass

from app.api.acl import ACL_CLAUSE, doc_acl

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass
class EditableDocument:
    document_id: str
    title: str
    paragraphs: list[str]
    source: str                  # "docx" | "text"
    docx: bytes | None = None
    version_id: str | None = None


def _original_uri(conn, document_id: str) -> tuple[str | None, str | None]:
    row = conn.execute(
        """
        SELECT coalesce(v.storage_uri, d.source_uri) AS uri, coalesce(v.mime_type, d.mime_type) AS mime
        FROM documents d LEFT JOIN document_versions v ON v.version_id = d.current_version_id
        WHERE d.document_id = %s
        """,
        (document_id,),
    ).fetchone()
    return (row["uri"], row["mime"]) if row else (None, None)


def load_editable(conn, document_id: str, member_id: str | None) -> EditableDocument | None:
    """None when the document does not exist or the member may not see it."""
    from app.chat.tools.document_tools import fetch_document_pages
    from app.storage.object_store import get_object_store

    row = conn.execute(
        f"""
        SELECT d.document_id, d.title, d.current_version_id
        FROM documents d LEFT JOIN permissions p ON p.matter_id = d.matter_id
        WHERE d.document_id = %(id)s AND {ACL_CLAUSE} AND {doc_acl('d')}
        """,
        {"id": document_id, "member_id": member_id},
    ).fetchone()
    if not row:
        return None
    uri, mime = _original_uri(conn, document_id)
    if uri and (mime == DOCX_MIME or str(uri).lower().endswith(".docx")):
        try:
            from docx import Document

            data = get_object_store().get(uri)
            paras = [p.text for p in Document(io.BytesIO(data)).paragraphs]
            return EditableDocument(document_id, row["title"], paras, "docx", data, row["current_version_id"])
        except Exception:
            pass  # unreadable original: fall back to the extracted text
    pages = fetch_document_pages(conn, document_id) or []
    paras = [" ".join(p.split()) for _page, text in pages for p in re.split(r"\n\s*\n", text) if p.strip()]
    return EditableDocument(document_id, row["title"], paras, "text", None, row["current_version_id"])
