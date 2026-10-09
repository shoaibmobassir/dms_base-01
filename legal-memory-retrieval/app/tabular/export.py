"""A tabular review as an Excel workbook: the table as the person sees it, and every answer's source (plan 22, W4.6)."""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEAD = PatternFill(start_color="3B1F2B", end_color="3B1F2B", fill_type="solid")
_HEAD_FONT = Font(bold=True, color="FFFFFF")
_WRAP = Alignment(vertical="top", wrap_text=True)
_STATUS = {"not_found": "Not found", "failed": "Could not answer", "pending": "Not run yet", "running": "Running",
           "stale": "Question changed — run again"}


def _cell_text(c: dict[str, Any]) -> str:
    if c.get("restricted"):
        return "Restricted"
    if c["status"] in ("done",) or c.get("edited"):
        return c.get("answer") or ""
    return _STATUS.get(c["status"], c["status"])


def workbook(view: dict[str, Any]) -> bytes:
    """``view`` is ``app.tabular.get_review`` for the person exporting (already redacted for them)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    cols = view["columns"]
    cells = {(c["row_id"], c["column_id"]): c for c in view["cells"]}
    headers = ["Document" if view["group_by"] == "document" else "Folder", "Document ID"] + [c["label"] for c in cols]
    ws.append(headers)
    for i in range(1, len(headers) + 1):
        h = ws.cell(row=1, column=i)
        h.fill, h.font, h.alignment = _HEAD, _HEAD_FONT, _WRAP
    for r in view["rows"]:
        label = "Restricted document" if r["restricted"] else (r["title"] or r["folder_path"] or "")
        values = [label, r["document_id"] or ""]
        values += [_cell_text(cells.get((r["row_id"], c["column_id"]), {"status": "pending"})) for c in cols]
        ws.append(values)
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 40 if i > 2 else (36 if i == 1 else 18)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = _WRAP
    ws.freeze_panes = "C2"

    src = wb.create_sheet("Sources")
    src_head = ["Row", "Question", "Answer", "Source document", "Page", "Quoted passage", "Quote found in document",
                "Answered by"]
    src.append(src_head)
    for i in range(1, len(src_head) + 1):
        h = src.cell(row=1, column=i)
        h.fill, h.font, h.alignment = _HEAD, _HEAD_FONT, _WRAP
    rows = {r["row_id"]: r for r in view["rows"]}
    labels = {c["column_id"]: c["label"] for c in cols}
    for c in view["cells"]:
        if c.get("restricted") or c["column_id"] not in labels:
            continue
        r = rows.get(c["row_id"])
        if not r or r["restricted"]:
            continue
        for cit in c.get("citations") or [{}]:
            if not cit and not c.get("edited"):
                continue
            src.append([r["title"] or r["folder_path"], labels[c["column_id"]], _cell_text(c), (cit or {}).get("document_id"),
                        (cit or {}).get("page"), (cit or {}).get("quote"), "yes" if (cit or {}).get("verified") else "no",
                        "a lawyer" if c.get("edited") else "the model"])
    for i, w in enumerate((36, 24, 40, 18, 6, 60, 10, 12), 1):
        src.column_dimensions[get_column_letter(i)].width = w
    for row in src.iter_rows(min_row=2):
        for c in row:
            c.alignment = _WRAP
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
