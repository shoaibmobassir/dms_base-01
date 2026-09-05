"""DOCX text extractor using python-docx (MIT licensed) — paragraph → pseudo-pages."""
from __future__ import annotations

import io
import tempfile
from pathlib import Path

from app.ingest.extractors.types import ExtractedDocument, join_pages

# Approximate page break every N characters for DOCX (no fixed pages in OOXML flow)
_CHARS_PER_PAGE = 3000


class DocxExtractor:
    """Extract structured text from DOCX, preserving paragraphs and headings."""

    def extract(self, path: str) -> tuple[str, int]:
        doc = self.extract_structured(path)
        return doc.text, doc.page_count

    def extract_structured(self, path: str) -> ExtractedDocument:
        from docx import Document

        document = Document(path)
        paragraphs: list[str] = []
        for p in document.paragraphs:
            text = (p.text or "").strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style.lower().startswith("heading"):
                # Normalize toward legal heading patterns for the block parser
                paragraphs.append(text if text.upper().startswith(
                    ("ARTICLE", "SECTION", "CLAUSE", "SCHEDULE", "EXHIBIT", "PART")
                ) else f"SECTION {text}")
            else:
                paragraphs.append(text)

        # Tables as pipe-separated blocks
        for table in document.tables:
            rows = []
            for row in table.rows:
                cells = [" ".join(c.text.split()) for c in row.cells]
                rows.append(" | ".join(cells))
            if rows:
                paragraphs.append("\n".join(rows))

        body = "\n\n".join(paragraphs)
        # Slice into pseudo-pages for Page → Block hierarchy
        page_texts: list[str] = []
        if not body:
            page_texts = [""]
        else:
            start = 0
            while start < len(body):
                end = min(len(body), start + _CHARS_PER_PAGE)
                if end < len(body):
                    cut = body.rfind("\n\n", start, end)
                    if cut > start + _CHARS_PER_PAGE // 2:
                        end = cut
                page_texts.append(body[start:end].strip())
                start = end
            page_texts = [p for p in page_texts if p] or [body]

        extracted = join_pages(page_texts)
        extracted.mime_type = (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        extracted.source_format = "docx"
        return extracted

    def extract_bytes(self, data: bytes) -> ExtractedDocument:
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=True) as tmp:
            tmp.write(data)
            tmp.flush()
            return self.extract_structured(tmp.name)
