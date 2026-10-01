"""Write edits into the lawyer's ORIGINAL .docx as Word tracked changes (plan 15, Part D3).

The existing redline export rebuilds a fresh document from plain text, which loses the
original formatting. Here the original file is kept and only the edited paragraphs change:

  replace       word-level diff of the paragraph's old and new text; unchanged words keep
                their runs (and formatting), removed words become <w:del>, new words <w:ins>
                with the formatting of the run they are inserted into
  delete        every run of the paragraph inside <w:del>, and the paragraph mark deleted
  insert_after  a new paragraph after the anchor, copying its paragraph properties and the
                formatting of its last run, all inside <w:ins> (paragraph mark inserted too)

Paragraph ids are indexes into ``Document.paragraphs`` (body paragraphs, tables excluded).
A paragraph whose runs hold more than text (fields, tabs, hyperlinks, drawings) is replaced
as a whole — deleted and re-inserted — rather than diffed, so nothing is silently dropped.
"""
from __future__ import annotations

import copy
import difflib
import io
import re
from datetime import datetime, timezone
from typing import Any

from docx import Document
from docx.oxml.ns import qn

_SIMPLE_CHILDREN = {qn("w:rPr"), qn("w:t")}
_TOKEN = re.compile(r"\s+|[^\s]+")


class _Ids:
    def __init__(self, start: int = 9000):
        self.n = start

    def next(self) -> str:
        self.n += 1
        return str(self.n)


def _mark(tag: str, ids: _Ids, author: str, date: str):
    from docx.oxml import OxmlElement

    el = OxmlElement(tag)
    el.set(qn("w:id"), ids.next())
    el.set(qn("w:author"), author)
    el.set(qn("w:date"), date)
    return el


def _run(text: str, rpr, deleted: bool = False):
    from docx.oxml import OxmlElement

    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    t = OxmlElement("w:delText" if deleted else "w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    return r


def _runs(p) -> list:
    return [c for c in p._p if c.tag == qn("w:r")]


def _is_simple(p) -> bool:
    kids = [c for c in p._p if c.tag != qn("w:pPr")]
    return all(c.tag == qn("w:r") and {g.tag for g in c} <= _SIMPLE_CHILDREN for c in kids)


def _segments(p) -> list[tuple[int, int, Any]]:
    """(start, end, rPr) for each run, over the paragraph's text."""
    out, pos = [], 0
    for r in _runs(p):
        text = "".join(t.text or "" for t in r.iter(qn("w:t")))
        out.append((pos, pos + len(text), r.find(qn("w:rPr"))))
        pos += len(text)
    return out


def _slices(segs, a: int, b: int):
    """Split [a, b) into pieces that each lie in one run: (start, end, rPr)."""
    for s, e, rpr in segs:
        lo, hi = max(a, s), min(b, e)
        if lo < hi:
            yield lo, hi, rpr


def _rpr_at(segs, pos: int):
    for s, e, rpr in segs:
        if s <= pos < e or (pos == e and pos > s):
            return rpr
    return segs[-1][2] if segs else None


def _clear_content(p) -> None:
    for c in list(p._p):
        if c.tag != qn("w:pPr"):
            p._p.remove(c)


def _mark_paragraph(p, tag: str, ids: _Ids, author: str, date: str) -> None:
    """Record the paragraph mark itself as inserted or deleted (pPr/rPr/w:ins|w:del)."""
    from docx.oxml import OxmlElement

    ppr = p._p.get_or_add_pPr()
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        ppr.append(rpr)
    rpr.append(_mark(tag, ids, author, date))


def _replace(p, new_text: str, ids: _Ids, author: str, date: str) -> bool:
    """Diff in place; returns False when the paragraph had to be replaced whole."""
    old_text = p.text
    segs = _segments(p)
    if not _is_simple(p) or not segs:
        _delete(p, ids, author, date, keep_mark=True)
        ins = _mark("w:ins", ids, author, date)
        ins.append(_run(new_text, segs[-1][2] if segs else None))
        p._p.append(ins)
        return False
    a_tok, b_tok = _TOKEN.findall(old_text), _TOKEN.findall(new_text)
    a_pos = [0]
    for tok in a_tok:
        a_pos.append(a_pos[-1] + len(tok))
    pieces = []
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a_tok, b_tok, autojunk=False).get_opcodes():
        a, b = a_pos[i1], a_pos[i2]
        if tag == "equal":
            pieces += [_run(old_text[lo:hi], rpr) for lo, hi, rpr in _slices(segs, a, b)]
            continue
        if tag in ("delete", "replace"):
            d = _mark("w:del", ids, author, date)
            for lo, hi, rpr in _slices(segs, a, b):
                d.append(_run(old_text[lo:hi], rpr, deleted=True))
            pieces.append(d)
        if tag in ("insert", "replace"):
            ins = _mark("w:ins", ids, author, date)
            ins.append(_run("".join(b_tok[j1:j2]), _rpr_at(segs, a)))
            pieces.append(ins)
    _clear_content(p)
    for piece in pieces:
        p._p.append(piece)
    return True


def _delete(p, ids: _Ids, author: str, date: str, keep_mark: bool = False) -> None:
    d = _mark("w:del", ids, author, date)
    for c in [c for c in p._p if c.tag != qn("w:pPr")]:
        p._p.remove(c)
        if c.tag == qn("w:r"):
            for t in list(c.iter(qn("w:t"))):
                t.tag = qn("w:delText")
        d.append(c)
    p._p.append(d)
    if not keep_mark:
        _mark_paragraph(p, "w:del", ids, author, date)


def _insert_after(p, text: str, ids: _Ids, author: str, date: str):
    new = copy.deepcopy(p._p)
    for c in [c for c in new if c.tag != qn("w:pPr")]:
        new.remove(c)
    ppr = new.find(qn("w:pPr"))
    if ppr is not None and ppr.find(qn("w:rPr")) is not None:
        ppr.remove(ppr.find(qn("w:rPr")))
    p._p.addnext(new)
    from docx.text.paragraph import Paragraph

    para = Paragraph(new, p._parent)
    segs = _segments(p)
    ins = _mark("w:ins", ids, author, date)
    ins.append(_run(text, segs[-1][2] if segs else None))
    new.append(ins)
    _mark_paragraph(para, "w:ins", ids, author, date)
    return para


def apply_tracked_changes(source: bytes, ops: list[dict], *, author: str = "Precentis Assistant",
                          date: str | None = None) -> tuple[bytes, dict]:
    """Apply paragraph-id ops (see module docstring) to a .docx; returns (new .docx bytes, stats)."""
    doc = Document(io.BytesIO(source))
    paras = list(doc.paragraphs)  # ids refer to the original paragraphs, fixed before any insertion
    when = date or datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    ids = _Ids()
    stats = {"replaced": 0, "replaced_whole": 0, "deleted": 0, "inserted": 0, "rejected": 0}
    inserts: dict[int, list[str]] = {}
    for op in ops:
        pid = op.get("pid")
        if not isinstance(pid, int) or not 0 <= pid < len(paras):
            stats["rejected"] += 1
            continue
        kind = op.get("op")
        if kind == "replace":
            if _replace(paras[pid], str(op.get("text") or ""), ids, author, when):
                stats["replaced"] += 1
            else:
                stats["replaced_whole"] += 1
        elif kind == "delete":
            _delete(paras[pid], ids, author, when)
            stats["deleted"] += 1
        elif kind == "insert_after":
            inserts.setdefault(pid, []).append(str(op.get("text") or ""))
        else:
            stats["rejected"] += 1
    for pid, texts in inserts.items():
        anchor = paras[pid]
        for text in texts:
            anchor = _insert_after(anchor, text, ids, author, when)
            stats["inserted"] += 1
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), stats


# ---------------------------------------------------------------------------
# Reading tracked changes back (validation: accept-all / reject-all views)
# ---------------------------------------------------------------------------

def _para_view(p_el, accept: bool) -> str | None:
    """Text of one w:p with all changes accepted (or rejected); None if the paragraph disappears."""
    mark = p_el.find(f"{qn('w:pPr')}/{qn('w:rPr')}")
    mark_del = mark is not None and mark.find(qn("w:del")) is not None
    mark_ins = mark is not None and mark.find(qn("w:ins")) is not None
    if (accept and mark_del) or (not accept and mark_ins):
        return None
    parts: list[str] = []
    for el in p_el.iter():
        if el.tag == qn("w:t"):
            inside_ins = any(a.tag == qn("w:ins") for a in el.iterancestors())
            if accept or not inside_ins:
                parts.append(el.text or "")
        elif el.tag == qn("w:delText") and not accept:
            parts.append(el.text or "")
    return "".join(parts)


def view(docx_bytes: bytes, accept: bool) -> list[str]:
    doc = Document(io.BytesIO(docx_bytes))
    out = []
    for p in doc.element.body.iter(qn("w:p")):
        if any(a.tag == qn("w:tbl") for a in p.iterancestors()):
            continue
        text = _para_view(p, accept)
        if text is not None:
            out.append(text)
    return out


def formatting_signature(p_el) -> tuple:
    """Style + per-run formatting and text of a paragraph (for 'untouched paragraphs are unchanged')."""
    ppr = p_el.find(qn("w:pPr"))
    style = ppr.find(qn("w:pStyle")).get(qn("w:val")) if ppr is not None and ppr.find(qn("w:pStyle")) is not None else ""
    runs = []
    for r in p_el.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        props = tuple(sorted((c.tag.split("}")[1], c.get(qn("w:val")) or "1") for c in rpr)) if rpr is not None else ()
        runs.append((props, "".join(t.text or "" for t in r.iter(qn("w:t")))))
    # Adjacent runs with identical formatting are one visual run; compare them merged.
    merged: list[list] = []
    for props, text in runs:
        if merged and merged[-1][0] == props:
            merged[-1][1] += text
        else:
            merged.append([props, text])
    return style, tuple((p, t) for p, t in merged)


def accept_all(docx_bytes: bytes) -> bytes:
    """The same file with every tracked change accepted: the clean text of the new version,
    original formatting kept (the stored original of that version)."""
    doc = Document(io.BytesIO(docx_bytes))
    body = doc.element.body
    for p in list(body.iter(qn("w:p"))):
        mark = p.find(f"{qn('w:pPr')}/{qn('w:rPr')}")
        if mark is not None and mark.find(qn("w:del")) is not None:
            p.getparent().remove(p)
            continue
        if mark is not None:
            for m in mark.findall(qn("w:ins")):
                mark.remove(m)
    for d in list(body.iter(qn("w:del"))):
        d.getparent().remove(d)
    for ins in list(body.iter(qn("w:ins"))):
        parent = ins.getparent()
        at = parent.index(ins)
        for child in list(ins):
            parent.insert(at, child)
            at += 1
        parent.remove(ins)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
