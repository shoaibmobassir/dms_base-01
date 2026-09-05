"""
DOCX Redline Generator: Produces native Microsoft Word OpenXML Tracked Changes (<w:ins>, <w:del>).
Clean-room independent implementation.
"""

from datetime import datetime
import io
import os
from typing import List, Optional, Tuple
import difflib
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


class DocxRedlineGenerator:
    """Injects native Word tracked revisions into .docx files for formal counterparty negotiation."""

    def __init__(self, author: str = "FirmOS AI Review"):
        self.author = author

    def create_tracked_diff_docx(
        self,
        original_text: str,
        revised_text: str,
        title: str = "Redlined Document",
    ) -> bytes:
        """Generates a .docx document where differences are marked as native Word tracked changes."""
        doc = docx.Document()
        doc.add_heading(title, level=1)

        # Tokenize by words/paragraphs for legal granular diff
        orig_paras = original_text.split("\n")
        rev_paras = revised_text.split("\n")

        matcher = difflib.SequenceMatcher(None, orig_paras, rev_paras)
        revision_id = 1
        now_str = datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ")

        for tag, i1, i2, j1, j2 in matcher.get_opcodes():
            if tag == "equal":
                for p_text in orig_paras[i1:i2]:
                    if p_text.strip():
                        doc.add_paragraph(p_text)
            elif tag == "replace":
                # Paragraph-level or word-level replacement
                orig_block = "\n".join(orig_paras[i1:i2])
                rev_block = "\n".join(rev_paras[j1:j2])
                p = doc.add_paragraph()
                revision_id = self._append_inline_tracked_diff(
                    p, orig_block, rev_block, revision_id, now_str
                )
            elif tag == "delete":
                for p_text in orig_paras[i1:i2]:
                    if p_text.strip():
                        p = doc.add_paragraph()
                        self._append_deletion(p, p_text, revision_id, now_str)
                        revision_id += 1
            elif tag == "insert":
                for p_text in rev_paras[j1:j2]:
                    if p_text.strip():
                        p = doc.add_paragraph()
                        self._append_insertion(p, p_text, revision_id, now_str)
                        revision_id += 1

        out = io.BytesIO()
        doc.save(out)
        return out.getvalue()

    def _append_inline_tracked_diff(
        self,
        paragraph,
        orig_text: str,
        rev_text: str,
        start_rev_id: int,
        date_str: str,
    ) -> int:
        """Appends word-level tracked insertions and deletions into a paragraph."""
        orig_words = orig_text.split(" ")
        rev_words = rev_text.split(" ")
        matcher = difflib.SequenceMatcher(None, orig_words, rev_words)
        rev_id = start_rev_id

        for tag, i1, i2, j1, j2 in matcher.get_opcodes():
            if tag == "equal":
                paragraph.add_run(" ".join(orig_words[i1:i2]) + " ")
            elif tag == "delete":
                del_words = " ".join(orig_words[i1:i2]) + " "
                self._append_deletion(paragraph, del_words, rev_id, date_str)
                rev_id += 1
            elif tag == "insert":
                ins_words = " ".join(rev_words[j1:j2]) + " "
                self._append_insertion(paragraph, ins_words, rev_id, date_str)
                rev_id += 1
            elif tag == "replace":
                del_words = " ".join(orig_words[i1:i2]) + " "
                ins_words = " ".join(rev_words[j1:j2]) + " "
                self._append_deletion(paragraph, del_words, rev_id, date_str)
                rev_id += 1
                self._append_insertion(paragraph, ins_words, rev_id, date_str)
                rev_id += 1

        return rev_id

    def _append_insertion(self, paragraph, text: str, rev_id: int, date_str: str):
        """Creates a <w:ins> OpenXML element with inner run and text."""
        ins_elm = parse_xml(
            f'<w:ins {nsdecls("w")} w:id="{rev_id}" w:author="{self.author}" w:date="{date_str}">'
            f'  <w:r>'
            f'    <w:rPr><w:color w:val="008000"/><w:u w:val="single"/></w:rPr>'
            f'    <w:t>{self._escape_xml(text)}</w:t>'
            f'  </w:r>'
            f'</w:ins>'
        )
        paragraph._p.append(ins_elm)

    def _append_deletion(self, paragraph, text: str, rev_id: int, date_str: str):
        """Creates a <w:del> OpenXML element with inner run and delText."""
        del_elm = parse_xml(
            f'<w:del {nsdecls("w")} w:id="{rev_id}" w:author="{self.author}" w:date="{date_str}">'
            f'  <w:r>'
            f'    <w:rPr><w:color w:val="FF0000"/><w:strike/></w:rPr>'
            f'    <w:delText>{self._escape_xml(text)}</w:delText>'
            f'  </w:r>'
            f'</w:del>'
        )
        paragraph._p.append(del_elm)

    @staticmethod
    def _escape_xml(text: str) -> str:
        return (
            text.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace('"', "&quot;")
            .replace("'", "&apos;")
        )
