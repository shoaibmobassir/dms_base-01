"""
Generation tools: generate_docx and generate_excel callable by the LLM agent.
Clean-room independent implementation for FirmOS legal assistant chatbot.
"""

from __future__ import annotations

import io
import os
import uuid
from pathlib import Path
from typing import Any

from app.config import settings
from app.db.connection import connect
from app.storage.object_store import sanitize_segment

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def generated_dir() -> Path:
    return (Path(settings.object_store_root) / "generated").resolve()


def _safe_filename(title: str, ext: str) -> str:
    # Titles come from the model (and so, indirectly, from document text): never
    # let them carry path separators or traversal segments into a filesystem path.
    flat = title.replace("/", " ").replace("\\", " ").replace(" ", "_")
    stem = sanitize_segment(flat).replace("/", "_")[:120] or "document"
    return f"{stem}.{ext}"


def _store(save: Any, title: str, ext: str, mime: str, owner_member_id: str | None) -> dict[str, Any]:
    """Save a generated file under generated/<artifact_id>/ and record its owner."""
    artifact_id = str(uuid.uuid4())
    filename = _safe_filename(title, ext)
    folder = generated_dir() / artifact_id
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / filename
    save(str(path))
    with connect() as conn:
        conn.execute(
            """
            INSERT INTO generated_artifacts
                (artifact_id, owner_member_id, filename, storage_path, mime_type, size_bytes)
            VALUES (%s, %s, %s, %s, %s, %s)
            """,
            (artifact_id, owner_member_id, filename, str(path), mime, path.stat().st_size),
        )
        conn.commit()
    download_url = f"/api/documents/{artifact_id}/download"
    return {
        "filename": filename,
        "download_url": download_url,
        "document_id": artifact_id,
        "version_id": str(uuid.uuid4()),
        "event": {
            "type": "doc_created",
            "filename": filename,
            "download_url": download_url,
            "document_id": artifact_id,
        },
    }


def generate_docx(title: str, sections: list[dict[str, Any]], owner_member_id: str | None = None) -> dict[str, Any]:
    """
    Generate a Word (.docx) document from structured section content.
    Returns a download URL and document metadata.
    """
    try:
        from docx import Document as DocxDocument
        from docx.shared import Pt, Inches
    except ImportError:
        return {"error": "python-docx not installed. Cannot generate DOCX files."}

    doc = DocxDocument()
    # Title
    doc.add_heading(title, level=0)

    for section in sections:
        page_break = section.get("pageBreak", False)
        if page_break:
            doc.add_page_break()

        heading = section.get("heading")
        level = section.get("level", 1)
        if heading:
            doc.add_heading(heading, level=min(level, 3))

        content = section.get("content", "")
        if content:
            for para_text in content.split("\n\n"):
                para_text = para_text.strip()
                if para_text:
                    doc.add_paragraph(para_text)

        # Tables
        table_data = section.get("table")
        if table_data:
            headers = table_data.get("headers", [])
            rows = table_data.get("rows", [])
            if headers:
                table = doc.add_table(rows=1 + len(rows), cols=len(headers))
                table.style = "Light Grid Accent 1"
                for i, header in enumerate(headers):
                    table.rows[0].cells[i].text = header
                for row_idx, row in enumerate(rows):
                    for col_idx, cell_val in enumerate(row):
                        if col_idx < len(headers):
                            table.rows[row_idx + 1].cells[col_idx].text = str(cell_val)

    return _store(doc.save, title, "docx", DOCX_MIME, owner_member_id)


def generate_excel(title: str, sheets: list[dict[str, Any]], owner_member_id: str | None = None) -> dict[str, Any]:
    """
    Generate an Excel (.xlsx) workbook from structured sheet data.
    Returns a download URL and document metadata.
    """
    try:
        import openpyxl
    except ImportError:
        return {"error": "openpyxl not installed. Cannot generate Excel files."}

    wb = openpyxl.Workbook()
    # Remove default sheet
    if wb.active:
        wb.remove(wb.active)

    for sheet_data in sheets:
        ws = wb.create_sheet(title=sheet_data.get("name", "Sheet"))
        columns = sheet_data.get("columns", [])
        rows = sheet_data.get("rows", [])

        # Header row
        for col_idx, col_name in enumerate(columns, 1):
            ws.cell(row=1, column=col_idx, value=col_name)

        # Data rows
        for row_idx, row in enumerate(rows, 2):
            for col_idx, cell_val in enumerate(row, 1):
                ws.cell(row=row_idx, column=col_idx, value=cell_val)

    return _store(wb.save, title, "xlsx", XLSX_MIME, owner_member_id)
